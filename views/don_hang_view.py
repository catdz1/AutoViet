"""
views/don_hang_view.py — Màn hình quản lý Đơn hàng
+ Form tạo đơn: Khách hàng mới/cũ + NV tự động điền
+ NV đăng nhập → tự điền tên mình vào đơn hàng
Tách từ views/other_views.py để mỗi màn hình nghiệp vụ nằm 1 file riêng.
"""
import os                  # ← THÊM DÒNG NÀY
import pandas as pd
import calendar, random
from datetime import datetime, date
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QTableWidget, QTableWidgetItem,
    QPushButton, QLineEdit, QDialog, QFormLayout, QComboBox,
    QMessageBox, QHeaderView, QLabel, QDoubleSpinBox, QTextEdit,
    QFrame, QApplication, QTabWidget, QScrollArea, QGridLayout, QInputDialog
)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor, QFont, QPixmap   # ← thêm QPixmap vào đây luôn
from database import get_conn

from views.base_widgets import BaseDialog, BaseView

# ════════════════════════════════════════════════════════════════
# ĐƠN HÀNG — KH mới/cũ + NV tự động
# ════════════════════════════════════════════════════════════════
class DonHangView(BaseView):
    COLS=[("MÃ ĐƠN","ma_don",False),("XE","ten_xe",True),("KHÁCH HÀNG","ten_kh",False),
          ("NHÂN VIÊN","ten_nv",False),("GIÁ BÁN","gia_ban_thuc",False),
          ("TRẠNG THÁI","trang_thai",False),("NGÀY ĐẶT","ngay_dat",False)]
    def __init__(self,current_user=None): super().__init__("page_don_hang",current_user)

    def _build_toolbar(self):
        # ── Nút Tạo đơn — xanh đậm nổi bật ──────────────
        b_add = QPushButton("➕  Tạo đơn")
        b_add.setObjectName("btn_add")
        b_add.setCursor(Qt.CursorShape.PointingHandCursor)
        b_add.setStyleSheet("""
            QPushButton{background:#2563eb;color:#ffffff;border:none;
                border-radius:8px;padding:8px 16px;font-size:13px;font-weight:700;}
            QPushButton:hover{background:#1d4ed8;}
            QPushButton:pressed{background:#1e40af;}
        """)
        b_add.clicked.connect(self._add)
        self._tbh.addWidget(b_add)

        # ── Nút Chi tiết — xám nhạt ───────────────────────
        b_ct = QPushButton("📋  Chi tiết")
        b_ct.setCursor(Qt.CursorShape.PointingHandCursor)
        b_ct.setStyleSheet("""
            QPushButton{background:#f1f5f9;color:#334155;border:1px solid #e2e8f0;
                border-radius:8px;padding:8px 14px;font-size:13px;font-weight:600;}
            QPushButton:hover{background:#e2e8f0;color:#1e293b;}
        """)
        b_ct.clicked.connect(self._xem_chitiet)
        self._tbh.addWidget(b_ct)

        # ── Nút In hóa đơn — xanh nhạt ───────────────────
        b_hd = QPushButton("🧾  In hóa đơn")
        b_hd.setCursor(Qt.CursorShape.PointingHandCursor)
        b_hd.setStyleSheet("""
            QPushButton{background:#eff6ff;color:#2563eb;
                border:1px solid #bfdbfe;
                border-radius:8px;padding:8px 14px;
                font-size:13px;font-weight:700;}
            QPushButton:hover{background:#2563eb;color:#ffffff;border-color:#2563eb;}
        """)
        b_hd.clicked.connect(self._in_pdf)
        self._tbh.addWidget(b_hd)

        # ── Nút In hợp đồng — tím nổi bật ────────────────
        b_hop = QPushButton("📝  In hợp đồng")
        b_hop.setCursor(Qt.CursorShape.PointingHandCursor)
        b_hop.setStyleSheet("""
            QPushButton{background:#f5f3ff;color:#7c3aed;
                border:1px solid #ddd6fe;
                border-radius:8px;padding:8px 14px;
                font-size:13px;font-weight:700;}
            QPushButton:hover{background:#7c3aed;color:#ffffff;border-color:#7c3aed;}
        """)
        b_hop.clicked.connect(self._in_hop_dong)
        self._tbh.addWidget(b_hop)

        # ── Nút Cập nhật TT — cam ─────────────────────────
        b_tt = QPushButton("🔄  Cập nhật TT")
        b_tt.setCursor(Qt.CursorShape.PointingHandCursor)
        b_tt.setStyleSheet("""
            QPushButton{background:#fff7ed;color:#c2410c;
                border:1px solid #fed7aa;
                border-radius:8px;padding:8px 14px;
                font-size:13px;font-weight:600;}
            QPushButton:hover{background:#c2410c;color:#ffffff;border-color:#c2410c;}
        """)
        b_tt.clicked.connect(self._update_status)
        self._tbh.addWidget(b_tt)

        # ── Nút Excel — xanh lá ───────────────────────────
        b_xl = QPushButton("📊  Excel")
        b_xl.setCursor(Qt.CursorShape.PointingHandCursor)
        b_xl.setStyleSheet("""
            QPushButton{background:#f0fdf4;color:#16a34a;
                border:1px solid #bbf7d0;
                border-radius:8px;padding:8px 14px;
                font-size:13px;font-weight:600;}
            QPushButton:hover{background:#16a34a;color:#ffffff;border-color:#16a34a;}
        """)
        b_xl.clicked.connect(self._export)
        self._tbh.addWidget(b_xl)
        self._btn("📱 Hiện QR", None, self._hien_qr)  # ← thêm đây
        self._btn("🔍 Quét QR", None, self._quet_qr)
        self._search_box("🔍 Tìm đơn hàng...")

    def _hien_qr(self):
        if not self._check_sel("xem QR"): return
        row = next((r for r in self._rows if r["id"] == self._sel_id), None)
        if not row: return

        from qr_utils import tao_qr_don_hang, qr_to_pixmap
        qr_path = tao_qr_don_hang(row)

        dlg = QDialog(self)
        dlg.setWindowTitle(f"📱 QR Code — {row['ma_don']}")
        dlg.setStyleSheet("QDialog{background:#0f1f35;} QLabel{background:transparent;color:#ffffff;}")
        dlg.setFixedSize(420, 520)

        lv = QVBoxLayout(dlg)
        lv.setContentsMargins(24, 20, 24, 20)
        lv.setSpacing(12)

        # Tiêu đề
        t = QLabel(f"📋  {row['ma_don']}")
        t.setStyleSheet("font-size:18px;font-weight:800;color:#60a5fa;")
        t.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lv.addWidget(t)

        # Thông tin đơn
        info = QLabel(
            f"🚗  {row.get('ten_xe', '')}\n"
            f"👤  {row.get('ten_kh', '')}\n"
            f"💰  {row.get('gia_ban_thuc', 0) / 1e9:.2f} tỷ ₫\n"
            f"📅  {row.get('ngay_dat', '')}"
        )
        info.setStyleSheet("font-size:13px;color:#94a3b8;line-height:1.6;")
        info.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lv.addWidget(info)

        # Ảnh QR lớn
        qr_lbl = QLabel()
        qr_lbl.setPixmap(qr_to_pixmap(qr_path, 280))
        qr_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        qr_lbl.setStyleSheet(
            "background:#ffffff;border-radius:16px;"
            "padding:16px;border:3px solid #2563eb;")
        lv.addWidget(qr_lbl)

        # Hướng dẫn
        hint = QLabel("📱 Dùng điện thoại quét mã này\nđể xem thông tin đơn hàng")
        hint.setStyleSheet("font-size:12px;color:#fbbf24;font-weight:600;")
        hint.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lv.addWidget(hint)

        # Nút đóng
        bh = QHBoxLayout();
        bh.addStretch()
        btn_c = QPushButton("✖  Đóng")
        btn_c.setStyleSheet("""
            QPushButton{background:#1e3a5f;color:#60a5fa;border:none;
            border-radius:8px;padding:10px 24px;font-size:13px;font-weight:700;}
            QPushButton:hover{background:#2563eb;color:white;}
        """)
        btn_c.clicked.connect(dlg.reject)
        bh.addWidget(btn_c);
        bh.addStretch()
        lv.addLayout(bh)

        dlg.exec()

    def _load(self,q=""):
        conn=get_conn()
        sql="""SELECT dh.*,x.hang_xe||' '||x.dong_xe as ten_xe,
               kh.ho_ten as ten_kh,nv.ho_ten as ten_nv
               FROM don_hang dh JOIN xe x ON dh.xe_id=x.id
               JOIN khach_hang kh ON dh.kh_id=kh.id
               JOIN nhan_vien nv ON dh.nv_id=nv.id"""
        p=[]
        if self.current_user.get("role")=="nhanvien" and self.current_user.get("nv_id"):
            sql+=" WHERE dh.nv_id=?"; p.append(self.current_user["nv_id"])
            if q: sql+=" AND (dh.ma_don LIKE ? OR kh.ho_ten LIKE ?)"; p+=[f"%{q}%"]*2
        else:
            if q: sql+=" WHERE dh.ma_don LIKE ? OR kh.ho_ten LIKE ?"; p=[f"%{q}%"]*2
        self._rows=[dict(r) for r in conn.execute(sql+" ORDER BY dh.id DESC",p).fetchall()]
        conn.close(); self._render()
        STATUS={"Đã giao xe":"#059669","Đã thanh toán":"#2563eb",
                "Chờ xử lý":"#d97706","Huỷ":"#dc2626"}
        for r,row in enumerate(self._rows):
            gi=QTableWidgetItem(f"{row['gia_ban_thuc']/1e9:.2f} tỷ")
            gi.setForeground(QColor("#059669"))
            gi.setFont(QFont("Segoe UI",13,QFont.Weight.Bold))
            self.tbl.setItem(r,4,gi)
            si=QTableWidgetItem(row["trang_thai"])
            si.setForeground(QColor(STATUS.get(row["trang_thai"],"#94a3b8")))
            si.setFont(QFont("Segoe UI",12,QFont.Weight.Bold))
            self.tbl.setItem(r,5,si)

    def _add(self):
        if DonHangDialog(self, current_user=self.current_user).exec():
            self._load()
            mw=self.window()
            if hasattr(mw,"_refresh_status"): mw._refresh_status()

    def _xem_chitiet(self):
        if not self._check_sel("xem chi tiết"): return
        row=next((r for r in self._rows if r["id"]==self._sel_id),None)
        if row: ChiTietDonHangDialog(self,row).exec()

    def _in_pdf(self):
        if not self._check_sel("in hóa đơn"): return
        try:
            from invoice_pdf import in_hoa_don
            fname=in_hoa_don(self._sel_id)
            QMessageBox.information(self,"✅ In hóa đơn thành công!",
                f"📄 File đã lưu:\n{fname}")
        except Exception as e:
            QMessageBox.critical(self,"❌ Lỗi in hóa đơn",str(e))

    def _in_hop_dong(self):
        if not self._check_sel("in hợp đồng"): return
        try:
            from hop_dong_pdf import in_hop_dong
            fname=in_hop_dong(self._sel_id)
            QMessageBox.information(self,"✅ In hợp đồng thành công!",
                f"📝 File hợp đồng đã lưu:\n{fname}\n\n"
                f"Mở file để in hoặc gửi cho khách hàng.")
        except ImportError:
            QMessageBox.critical(self,"❌ Thiếu thư viện",
                "Cần cài reportlab:\npip install reportlab")
        except Exception as e:
            QMessageBox.critical(self,"❌ Lỗi in hợp đồng",str(e))

    def _update_status(self):
        if not self._check_sel("cập nhật"): return
        from PyQt6.QtWidgets import QInputDialog

        RANK = {"Chờ xử lý": 0, "Cho xu ly": 0, "Đã thanh toán": 1, "Đã giao xe": 2}

        conn = get_conn()
        cur = conn.execute("SELECT trang_thai FROM don_hang WHERE id=?", (self._sel_id,)).fetchone()
        cur_status = cur["trang_thai"] if cur else None

        if cur_status == "Huỷ":
            conn.close()
            QMessageBox.warning(self, "", "Đơn đã huỷ, không thể đổi sang trạng thái khác!")
            return

        status, ok = QInputDialog.getItem(self, "Cập nhật", "Trạng thái:",
            ["Chờ xử lý", "Đã thanh toán", "Đã giao xe", "Huỷ"], 0, False)
        if not ok:
            conn.close(); return

        # Chỉ được Huỷ khi đơn đang "Chờ xử lý"
        if status == "Huỷ" and cur_status not in ("Chờ xử lý", "Cho xu ly"):
            conn.close()
            QMessageBox.warning(self, "",
                "Chỉ có thể huỷ đơn khi đang ở trạng thái «Chờ xử lý»!\n"
                "Đơn này đã ở trạng thái khác, không thể huỷ trực tiếp.")
            return

        # Không cho chuyển NGƯỢC (vd: "Đã giao xe" -> "Chờ xử lý")
        if status != "Huỷ" and RANK.get(status, 0) < RANK.get(cur_status, 0):
            conn.close()
            QMessageBox.warning(self, "",
                f"Không thể chuyển ngược từ «{cur_status}» về «{status}»!")
            return

        conn.execute("UPDATE don_hang SET trang_thai=? WHERE id=?", (status, self._sel_id))
        if status == "Đã giao xe":
            dh = conn.execute("SELECT xe_id FROM don_hang WHERE id=?", (self._sel_id,)).fetchone()
            if dh: conn.execute("UPDATE xe SET trang_thai='Đã bán' WHERE id=?", (dh[0],))
        if status == "Huỷ":
            conn.execute("UPDATE don_hang SET trang_thai_tt='Huỷ' WHERE id=?", (self._sel_id,))
        conn.commit(); conn.close(); self._load()

    def _export(self):
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        import datetime

        conn = get_conn()
        rows = conn.execute("""
            SELECT dh.ma_don, x.hang_xe||' '||x.dong_xe, kh.ho_ten,
                   nv.ho_ten, dh.gia_ban_thuc, dh.trang_thai, dh.ngay_dat
            FROM don_hang dh
            JOIN xe x ON dh.xe_id=x.id
            JOIN khach_hang kh ON dh.kh_id=kh.id
            JOIN nhan_vien nv ON dh.nv_id=nv.id
            ORDER BY dh.id DESC
        """).fetchall()
        conn.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Đơn hàng"

        ws.merge_cells("A1:G1")
        ws["A1"] = "BÁO CÁO ĐƠN HÀNG — AUTOVIET"
        ws["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="059669")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        ws.merge_cells("A2:G2")
        ws["A2"] = f"Ngày xuất: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
        ws["A2"].alignment = Alignment(horizontal="right")
        ws.row_dimensions[2].height = 20

        headers = ["MÃ ĐƠN", "XE", "KHÁCH HÀNG", "NHÂN VIÊN", "GIÁ BÁN", "TRẠNG THÁI", "NGÀY ĐẶT"]
        col_widths = [12, 30, 22, 22, 16, 16, 14]
        thin = Side(style="thin", color="BBF7D0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for i, (h, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=3, column=i, value=h)
            cell.font = Font(name="Segoe UI", size=11, bold=True, color="065F46")
            cell.fill = PatternFill("solid", fgColor="D1FAE5")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[3].height = 28

        fill_white = PatternFill("solid", fgColor="FFFFFF")
        fill_alt   = PatternFill("solid", fgColor="F0FDF4")
        font_data  = Font(name="Segoe UI", size=11, color="1E293B")
        font_ma    = Font(name="Segoe UI", size=11, bold=True, color="059669")

        STATUS_COLOR = {
            "Đã giao xe": "059669",
            "Đã thanh toán": "2563EB",
            "Chờ xử lý": "D97706",
            "Huỷ": "DC2626",
        }

        for ri, row in enumerate(rows, 4):
            fill = fill_white if ri % 2 == 0 else fill_alt
            ws.row_dimensions[ri].height = 24
            for ci, val in enumerate(row, 1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.fill = fill
                cell.border = border
                cell.alignment = Alignment(vertical="center",
                    horizontal="center" if ci in (1, 5, 6, 7) else "left")
                if ci == 1:
                    cell.font = font_ma
                elif ci == 5:
                    cell.font = Font(name="Segoe UI", size=11, bold=True, color="059669")
                elif ci == 6:
                    color = STATUS_COLOR.get(str(val), "64748B")
                    cell.font = Font(name="Segoe UI", size=11, bold=True, color=color)
                else:
                    cell.font = font_data

        ws.freeze_panes = "A4"
        fname = f"BaoCao_DonHang_{datetime.datetime.now().strftime('%d%m%Y_%H%M')}.xlsx"
        wb.save(fname)
        QMessageBox.information(self, "Excel", f"✅ Đã xuất: {fname}")

    def _quet_qr(self):
        from qr_scanner import QRScannerDialog
        dlg = QRScannerDialog(self)
        dlg.don_hang_found.connect(self._tim_don_by_ma)
        dlg.exec()

    def _tim_don_by_ma(self, ma_don):
        # Highlight dòng trong bảng
        for r, row in enumerate(self._rows):
            if row.get("ma_don") == ma_don:
                self.tbl.selectRow(r)
                self._sel_id = row["id"]
                break

        # Load đầy đủ thông tin từ DB
        from database import get_conn
        conn = get_conn()
        dh = conn.execute("""
            SELECT dh.ma_don, x.hang_xe||' '||x.dong_xe, x.mau_sac,
                   x.nam_sx, x.so_khung, x.so_may,
                   kh.ho_ten, kh.so_dt, kh.dia_chi,
                   nv.ho_ten, dh.gia_ban_thuc, dh.chiet_khau,
                   dh.phuong_thuc, dh.trang_thai, dh.ngay_dat, dh.ghi_chu
            FROM don_hang dh
            JOIN xe x ON dh.xe_id=x.id
            JOIN khach_hang kh ON dh.kh_id=kh.id
            JOIN nhan_vien nv ON dh.nv_id=nv.id
            WHERE dh.ma_don=?
        """, (ma_don,)).fetchone()
        conn.close()

        if not dh:
            QMessageBox.warning(self, "Không tìm thấy",
                                f"Không có đơn hàng {ma_don}!");
            return

        # Dialog hiển thị đầy đủ
        dlg = QDialog(self)
        dlg.setWindowTitle(f"✅ Đơn hàng — {ma_don}")
        dlg.setMinimumSize(520, 600)
        dlg.setStyleSheet("""
            QDialog{background:#f0f4f8;}
            QLabel{background:transparent;color:#1e293b;}
            QPushButton{border-radius:8px;font-size:13px;font-weight:700;padding:10px 20px;}
        """)

        root = QVBoxLayout(dlg)
        root.setContentsMargins(0, 0, 0, 0);
        root.setSpacing(0)

        # Header navy
        hdr = QWidget()
        hdr.setStyleSheet("""background:qlineargradient(x1:0,y1:0,x2:1,y2:0,
            stop:0 #0f1f35,stop:1 #1e3a5f);""")
        hl = QVBoxLayout(hdr);
        hl.setContentsMargins(24, 20, 24, 20);
        hl.setSpacing(4)

        t1 = QLabel(f"✅  ĐƠN HÀNG {dh[0]}")
        t1.setStyleSheet("font-size:20px;font-weight:900;color:#ffffff;")
        t1.setAlignment(Qt.AlignmentFlag.AlignCenter)

        tt = dh[13]
        tt_color = {"Đã giao xe": "#4ade80", "Đã thanh toán": "#60a5fa",
                    "Chờ xử lý": "#fbbf24", "Huỷ": "#f87171"}.get(tt, "#94a3b8")
        t2 = QLabel(f"● {tt}")
        t2.setStyleSheet(f"font-size:14px;font-weight:700;color:{tt_color};")
        t2.setAlignment(Qt.AlignmentFlag.AlignCenter)

        hl.addWidget(t1);
        hl.addWidget(t2)
        root.addWidget(hdr)

        # Content
        scroll = QScrollArea();
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)
        content = QWidget();
        content.setStyleSheet("background:#f0f4f8;")
        cl = QVBoxLayout(content);
        cl.setContentsMargins(16, 16, 16, 16);
        cl.setSpacing(12)

        def section(title_text, color="#2563eb"):
            sec = QWidget()
            sec.setStyleSheet(f"background:#ffffff;border-radius:12px;border:1px solid #e2e8f0;")
            sl = QVBoxLayout(sec);
            sl.setContentsMargins(16, 14, 16, 14);
            sl.setSpacing(8)
            ttl = QLabel(title_text)
            ttl.setStyleSheet(f"font-size:13px;font-weight:800;color:#0f1f35;"
                              f"border-left:4px solid {color};padding-left:8px;")
            sl.addWidget(ttl)
            return sec, sl

        def row_info(layout, label, value, val_color="#1e293b"):
            rw = QHBoxLayout()
            lk = QLabel(label);
            lk.setFixedWidth(150)
            lk.setStyleSheet("font-size:12px;color:#64748b;font-weight:600;")
            lv2 = QLabel(str(value) if value else "—")
            lv2.setStyleSheet(f"font-size:13px;font-weight:700;color:{val_color};")
            lv2.setWordWrap(True)
            rw.addWidget(lk);
            rw.addWidget(lv2, 1)
            layout.addLayout(rw)

        # Thông tin xe
        sec1, sl1 = section("🚗  THÔNG TIN XE", "#2563eb")
        row_info(sl1, "Tên xe:", dh[1])
        row_info(sl1, "Màu sắc:", dh[2])
        row_info(sl1, "Năm SX:", dh[3])
        row_info(sl1, "Số khung:", dh[4])
        row_info(sl1, "Số máy:", dh[5])
        cl.addWidget(sec1)

        # Thông tin khách hàng
        sec2, sl2 = section("👤  KHÁCH HÀNG", "#16a34a")
        row_info(sl2, "Họ tên:", dh[6])
        row_info(sl2, "Số ĐT:", dh[7], "#16a34a")
        row_info(sl2, "Địa chỉ:", dh[8])
        cl.addWidget(sec2)

        # Thông tin thanh toán
        gia = float(dh[10] or 0)
        ck = float(dh[11] or 0)
        sec3, sl3 = section("💰  THANH TOÁN", "#d97706")
        row_info(sl3, "Giá bán:", f"{gia / 1e9:.3f} tỷ ₫", "#16a34a")
        row_info(sl3, "Chiết khấu:", f"-{ck / 1e6:.0f}tr ₫", "#dc2626")
        row_info(sl3, "Thực thu:", f"{(gia - ck) / 1e9:.3f} tỷ ₫", "#2563eb")
        row_info(sl3, "Thanh toán:", dh[12])
        row_info(sl3, "Ngày đặt:", dh[14])
        cl.addWidget(sec3)

        # Nhân viên + ghi chú
        sec4, sl4 = section("🧑‍💼  NHÂN VIÊN BÁN", "#7c3aed")
        row_info(sl4, "NV phụ trách:", dh[9], "#7c3aed")
        row_info(sl4, "Ghi chú:", dh[15] or "—")
        cl.addWidget(sec4)
        cl.addStretch()

        scroll.setWidget(content)
        root.addWidget(scroll, 1)

        # Footer buttons
        ftr = QWidget()
        ftr.setStyleSheet("background:#ffffff;border-top:1px solid #e2e8f0;")
        fl = QHBoxLayout(ftr);
        fl.setContentsMargins(16, 12, 16, 12);
        fl.setSpacing(10)

        btn_pdf = QPushButton("🖨️ In hóa đơn")
        btn_pdf.setStyleSheet("background:#2563eb;color:white;border:none;")
        btn_pdf.clicked.connect(lambda: self._in_pdf_by_ma(ma_don, dlg))

        btn_close = QPushButton("✖ Đóng")
        btn_close.setStyleSheet("background:#f8fafc;color:#64748b;border:1px solid #e2e8f0;")
        btn_close.clicked.connect(dlg.reject)

        fl.addWidget(btn_pdf);
        fl.addStretch();
        fl.addWidget(btn_close)
        root.addWidget(ftr)
        dlg.exec()

    def _in_pdf_by_ma(self, ma_don, parent_dlg=None):
        """In PDF theo mã đơn"""
        conn = get_conn()
        dh = conn.execute("SELECT id FROM don_hang WHERE ma_don=?",
                          (ma_don,)).fetchone()
        conn.close()
        if not dh:
            QMessageBox.warning(self, "", "Không tìm thấy đơn!");
            return
        try:
            from invoice_pdf import in_hoa_don
            fname = in_hoa_don(dh[0])
            QMessageBox.information(self, "✅ In hóa đơn", f"Đã xuất: {fname}")
            import os;
            os.startfile(fname)
        except Exception as e:
            QMessageBox.critical(self, "Lỗi", str(e))

class ChiTietDonHangDialog(BaseDialog):
    def __init__(self,parent=None,row=None):
        super().__init__(parent,f"Chi tiết — {row['ma_don']}",580); self.row=row; self._build_ct()
    def _build_ct(self):
        conn=get_conn()
        dh=conn.execute("""SELECT dh.*,x.hang_xe||' '||x.dong_xe as ten_xe,x.ma_xe,
            x.so_khung,x.mau_sac,x.nam_sx,kh.ho_ten as ten_kh,kh.so_dt,kh.dia_chi,
            nv.ho_ten as ten_nv FROM don_hang dh JOIN xe x ON dh.xe_id=x.id
            JOIN khach_hang kh ON dh.kh_id=kh.id JOIN nhan_vien nv ON dh.nv_id=nv.id
            WHERE dh.id=?""",(self.row["id"],)).fetchone()
        conn.close()
        if not dh: return
        dh=dict(dh)
        self._add_title(f"📋  ĐƠN HÀNG {dh['ma_don']}")

        def row_lv(k, v, vc="#00274c"):
            rw = QHBoxLayout();
            lk = QLabel(k);
            lk.setFixedWidth(130)
            lk.setStyleSheet("color:#64748b;font-size:13px;font-weight:600;background:transparent;")
            lv = QLabel(str(v));
            lv.setStyleSheet(f"color:{vc};font-size:14px;background:transparent;font-weight:600;")
            lv.setWordWrap(True); rw.addWidget(lk); rw.addWidget(lv,1); return rw

        for k, v, c in [("Mã đơn", dh['ma_don'], "#1e40af"), ("Ngày đặt", dh['ngay_dat'], "#00274c"),
                        ("Phương thức", dh['phuong_thuc'], "#0891b2"),
                        ("Trạng thái", dh['trang_thai'],
                         {"Đã giao xe": "#065f46", "Đã thanh toán": "#1e40af", "Chờ xử lý": "#92400e",
                          "Huỷ": "#991b1b"}.get(dh['trang_thai'], "#64748b")),
                        ("Xe", dh['ten_xe'], "#00274c"), ("Khách hàng", dh['ten_kh'], "#00274c"),
                        ("Nhân viên BH", dh['ten_nv'], "#00274c"),
                        ("Giá bán", f"{dh['gia_ban_thuc']:,.0f} ₫", "#059669"),
                        ("Chiết khấu", f"- {dh.get('chiet_khau', 0) or 0:,.0f} ₫", "#dc2626"),
                        ("THỰC THU", f"{dh['gia_ban_thuc'] - (dh.get('chiet_khau', 0) or 0):,.0f} ₫", "#065f46")]:
            self._main_lv.addLayout(row_lv(k, v, c))
        bh=QHBoxLayout(); bh.addStretch()
        btn_pdf=QPushButton("📄 In hóa đơn PDF"); btn_pdf.setObjectName("btn_pdf")
        btn_pdf.clicked.connect(lambda: self._in_pdf(dh["id"]))
        bc=QPushButton("Đóng"); bc.clicked.connect(self.reject)
        bh.addWidget(btn_pdf); bh.addWidget(bc); self._main_lv.addLayout(bh)
    def _in_pdf(self,dh_id):
        try:
            from invoice_pdf import in_hoa_don
            fname=in_hoa_don(dh_id); QMessageBox.information(self,"OK",f"✅ {fname}")
        except Exception as e: QMessageBox.critical(self,"Lỗi",str(e))
    def _save(self): pass

class DonHangDialog(BaseDialog):
    def __init__(self, parent=None, current_user=None):
        super().__init__(parent, "Tạo đơn hàng mới", 580)
        self.current_user = current_user or {}
        self._add_title("📋  THÔNG TIN ĐƠN HÀNG")
        form = self._add_form()

        conn = get_conn()
        xe_rows = conn.execute(
            "SELECT id,ma_xe,hang_xe,dong_xe,gia_ban FROM xe WHERE trang_thai='Còn hàng'"
        ).fetchall()
        nv_rows = conn.execute(
            "SELECT id,ma_nv,ho_ten FROM nhan_vien WHERE trang_thai='Đang làm' ORDER BY ma_nv"
        ).fetchall()
        kh_rows = conn.execute(
            "SELECT id,ma_kh,ho_ten,so_dt FROM khach_hang ORDER BY id DESC"
        ).fetchall()
        cnt = conn.execute("SELECT COUNT(*) FROM don_hang").fetchone()[0]
        conn.close()

        # Mã đơn tự động — dò tìm mã chưa tồn tại để tránh trùng (UNIQUE constraint)
        # thay vì chỉ dùng COUNT(*)+1 (dễ bị trùng nếu đơn hàng từng bị xoá/tạo lệch số thứ tự)
        conn_ma = get_conn()
        so = cnt + 1
        ma_moi = f"DH{so:03d}"
        while conn_ma.execute("SELECT id FROM don_hang WHERE ma_don=?", (ma_moi,)).fetchone():
            so += 1
            ma_moi = f"DH{so:03d}"
        conn_ma.close()
        self.f_ma = QLineEdit(ma_moi)
        self.f_ma.setReadOnly(True)
        self.f_ma.setStyleSheet("color:#64748b;background:#f1f5f9;border-radius:8px;padding:6px;")

        # Xe
        self.f_xe = QComboBox()
        for r in xe_rows:
            self.f_xe.addItem(f"{r[1]} — {r[2]} {r[3]}  ({r[4]/1e9:.2f} tỷ)", r[0])
        self.f_xe.currentIndexChanged.connect(self._auto_fill_gia)

        # ── KHÁCH HÀNG: 2 tab ────────────────────────────────────────────
        kh_tabs = QTabWidget()
        kh_tabs.setStyleSheet("""
            QTabBar::tab{padding:6px 14px;font-size:11px;font-weight:600;
                color:#64748b;background:#13151c;border:none;
                border-bottom:2px solid transparent;}
            QTabBar::tab:selected{color:#a78bfa;border-bottom:2px solid #7c3aed;}
            QTabWidget::pane{border:1px solid #2c3050;border-radius:8px;}
        """)

        # Tab khách mới
        tab_new = QWidget()
        tnl = QVBoxLayout(tab_new); tnl.setContentsMargins(10,10,10,10); tnl.setSpacing(6)
        f2 = QFormLayout(); f2.setSpacing(8); f2.setLabelAlignment(Qt.AlignmentFlag.AlignRight)
        self.f_kh_ten   = QLineEdit(); self.f_kh_ten.setPlaceholderText("Họ và tên đầy đủ *")
        self.f_kh_sdt   = QLineEdit(); self.f_kh_sdt.setPlaceholderText("Số điện thoại *")
        self.f_kh_email = QLineEdit(); self.f_kh_email.setPlaceholderText("email@gmail.com")
        self.f_kh_cmnd  = QLineEdit(); self.f_kh_cmnd.setPlaceholderText("CMND / CCCD")
        self.f_kh_dc    = QLineEdit(); self.f_kh_dc.setPlaceholderText("Địa chỉ khách hàng")
        self.f_kh_loai  = QComboBox(); self.f_kh_loai.addItems(["Cá nhân","Doanh nghiệp"])
        for l,w in [("Họ tên *",self.f_kh_ten),("Số ĐT *",self.f_kh_sdt),
                    ("Email",self.f_kh_email),("CMND/CCCD",self.f_kh_cmnd),
                    ("Địa chỉ",self.f_kh_dc),("Loại KH",self.f_kh_loai)]:
            f2.addRow(l,w)
        tnl.addLayout(f2)
        kh_tabs.addTab(tab_new, "👤 Khách hàng mới")

        # Tab khách cũ
        tab_old = QWidget()
        tol = QVBoxLayout(tab_old); tol.setContentsMargins(10,10,10,10); tol.setSpacing(6)
        note = QLabel("Chọn khách hàng đã mua xe trước đây:")
        note.setStyleSheet("color:#64748b;font-size:12px;background:transparent;")
        self.f_kh_cu = QComboBox()
        for r in kh_rows:
            self.f_kh_cu.addItem(f"{r[1]} — {r[2]} ({r[3]})", r[0])
        tol.addWidget(note); tol.addWidget(self.f_kh_cu); tol.addStretch()
        kh_tabs.addTab(tab_old, "📋 Khách hàng cũ")
        self.kh_tabs = kh_tabs

        # ── NHÂN VIÊN: tự điền NV đang login ────────────────────────────
        self._auto_nv_id = self.current_user.get("nv_id")
        self._auto_nv_name = "—"

        if self._auto_nv_id:
            conn2 = get_conn()
            nv = conn2.execute(
                "SELECT ma_nv,ho_ten FROM nhan_vien WHERE id=?",
                (self._auto_nv_id,)
            ).fetchone()
            conn2.close()
            if nv: self._auto_nv_name = f"{nv[0]} — {nv[1]}"

        if self.current_user.get("role") == "admin" or not self._auto_nv_id:
            # Admin hoặc NV chưa liên kết → cho chọn
            self.f_nv = QComboBox()
            for r in nv_rows:
                self.f_nv.addItem(f"{r[1]} — {r[2]}", r[0])
            nv_widget = self.f_nv
        else:
            # NV → tự điền, không cho đổi
            self.f_nv = None
            nv_widget = QLabel(f"✅  {self._auto_nv_name}")
            nv_widget.setStyleSheet(
                "background:#f0fdf4;color:#065f46;border-radius:8px;"
                "border:0.5px solid #a7f3d0;"
                "padding:9px 12px;font-size:13px;font-weight:600;")

        # Giá bán + chiết khấu
        self.f_gia = QDoubleSpinBox()
        self.f_gia.setRange(0,1e11); self.f_gia.setSingleStep(1e6)
        self.f_gia.setDecimals(0); self.f_gia.setSuffix(" ₫")
        self.f_ck = QDoubleSpinBox()
        self.f_ck.setRange(0,1e10); self.f_ck.setSingleStep(500000)
        self.f_ck.setDecimals(0); self.f_ck.setSuffix(" ₫")
        self.f_tt = QComboBox()
        self.f_tt.addItems(["Tiền mặt","Chuyển khoản","Trả góp","Thẻ tín dụng","Vay NH"])
        self.f_ghi = QLineEdit(); self.f_ghi.setPlaceholderText("Ghi chú thêm...")

        for l,w in [
            ("Mã đơn",      self.f_ma),
            ("Xe *",        self.f_xe),
            ("Khách hàng *",kh_tabs),
            ("Nhân viên BH",nv_widget),
            ("Giá bán *",   self.f_gia),
            ("Chiết khấu",  self.f_ck),
            ("Thanh toán",  self.f_tt),
            ("Ghi chú",     self.f_ghi),
        ]: form.addRow(l,w)

        self._add_buttons("✅  Tạo đơn hàng")
        self._auto_fill_gia()

    def _auto_fill_gia(self):
        xe_id = self.f_xe.currentData()
        if xe_id:
            conn = get_conn()
            xe = conn.execute("SELECT gia_ban FROM xe WHERE id=?", (xe_id,)).fetchone()
            conn.close()
            if xe: self.f_gia.setValue(float(xe[0]))

    def _save(self):
        ma  = self.f_ma.text().strip()
        xe_id = self.f_xe.currentData()
        gia = self.f_gia.value()
        if not xe_id or gia <= 0:
            QMessageBox.warning(self,"","Chọn xe và nhập giá bán!"); return

        conn = get_conn()
        try:
            # Xử lý khách hàng
            if self.kh_tabs.currentIndex() == 0:
                ten = self.f_kh_ten.text().strip()
                sdt = self.f_kh_sdt.text().strip()
                if not ten or not sdt:
                    QMessageBox.warning(self,"","Nhập đủ Họ tên và Số ĐT khách hàng!"); return

                dup = conn.execute(
                    "SELECT ma_kh, ho_ten FROM khach_hang WHERE so_dt=?", (sdt,)
                ).fetchone()
                if dup:
                    QMessageBox.warning(self, "",
                        f"Số điện thoại «{sdt}» đã thuộc khách hàng {dup[0]} — {dup[1]}!\n"
                        f"Vui lòng chọn khách hàng cũ ở tab bên cạnh thay vì tạo mới.")
                    return

                cnt_kh = conn.execute("SELECT COUNT(*) FROM khach_hang").fetchone()[0]
                ma_kh = f"KH{cnt_kh+1:03d}"
                while conn.execute("SELECT id FROM khach_hang WHERE ma_kh=?",(ma_kh,)).fetchone():
                    cnt_kh+=1; ma_kh=f"KH{cnt_kh+1:03d}"
                c = conn.cursor()
                c.execute("""INSERT INTO khach_hang(ma_kh,ho_ten,so_dt,email,dia_chi,cmnd,loai_kh)
                    VALUES(?,?,?,?,?,?,?)""",
                    (ma_kh,ten,sdt,self.f_kh_email.text(),
                     self.f_kh_dc.text(),self.f_kh_cmnd.text(),
                     self.f_kh_loai.currentText()))
                kh_id = c.lastrowid
            else:
                kh_id = self.f_kh_cu.currentData()
                if not kh_id:
                    QMessageBox.warning(self,"","Chọn khách hàng!"); return

            # Xử lý NV
            nv_id = self.f_nv.currentData() if self.f_nv else self._auto_nv_id
            if not nv_id:
                QMessageBox.warning(self,"","Không xác định được nhân viên!"); return

            conn.execute("""INSERT INTO don_hang(ma_don, xe_id, kh_id, nv_id,
                                                 gia_ban_thuc, chiet_khau, phuong_thuc, ghi_chu,
                                                 trang_thai_tt, so_tien_da_tt)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                         (ma, xe_id, kh_id, nv_id, gia, self.f_ck.value(),
                          self.f_tt.currentText(), self.f_ghi.text(),
                          'Chưa thanh toán', 0))
            conn.commit()
            QMessageBox.information(self, "✅ Thành công!", f"Tạo đơn hàng {ma} thành công!")

            # Tạo QR code (tính năng phụ — nếu lỗi thì KHÔNG được làm người dùng
            # tưởng nhầm là tạo đơn hàng thất bại, vì đơn đã lưu thành công ở trên rồi)
            try:
                from qr_utils import tao_qr_don_hang
                from database import get_conn as _gc
                conn2 = _gc()
                dh = conn2.execute("""
                    SELECT dh.*, x.hang_xe||' '||x.dong_xe as ten_xe,
                           kh.ho_ten as ten_kh, nv.ho_ten as ten_nv
                    FROM don_hang dh
                    JOIN xe x ON dh.xe_id=x.id
                    JOIN khach_hang kh ON dh.kh_id=kh.id
                    JOIN nhan_vien nv ON dh.nv_id=nv.id
                    WHERE dh.ma_don=?
                """, (ma,)).fetchone()
                conn2.close()
                if dh:
                    tao_qr_don_hang(dict(dh))
            except ImportError:
                QMessageBox.information(self, "Thiếu thư viện tạo QR",
                    "Đơn hàng đã tạo thành công!\n\n"
                    "Không tạo được mã QR do máy chưa cài thư viện 'qrcode'.\n"
                    "Nếu muốn dùng tính năng QR, chạy lệnh:\n"
                    "pip install qrcode[pil]")
            except Exception as e:
                QMessageBox.information(self, "Không tạo được mã QR",
                    f"Đơn hàng đã tạo thành công!\nChỉ riêng mã QR bị lỗi: {e}")

            self.accept()
        except Exception as e:
            QMessageBox.critical(self, "Lỗi", str(e))
        finally:
            conn.close()
