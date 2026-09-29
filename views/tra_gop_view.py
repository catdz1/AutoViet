"""
views/tra_gop_view.py — Quản lý thanh toán trả góp
File MỚI — thêm vào views/
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QTableWidget, QTableWidgetItem, QHeaderView, QFrame,
    QDialog, QFormLayout, QComboBox, QLineEdit, QMessageBox,
    QDoubleSpinBox, QProgressBar, QScrollArea, QDateEdit, QFileDialog
)
from PyQt6.QtCore import Qt, QDate, QTimer
from PyQt6.QtGui import QColor, QFont
from database import get_conn
from datetime import datetime, date


import calendar

def tinh_tien_tra_gop(tong, lai_suat, so_ky):
    """
    Tính số tiền trả mỗi kỳ.
    
    lai_suat truyền vào theo %/tháng.
    Ví dụ: 0.8 nghĩa là 0.8%.
    """
    if so_ky <= 0:
        return 0

    lai = lai_suat / 100

    if lai > 0:
        return (
            tong * lai * (1 + lai) ** so_ky
            / ((1 + lai) ** so_ky - 1)
        )

    return tong / so_ky

def _add_months(d, n):
    """Cộng n tháng vào ngày d (cắt về cuối tháng nếu tháng đích ngắn hơn)."""
    m = d.month - 1 + n
    y = d.year + m // 12
    m = m % 12 + 1
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


def _ngay_ky_tiep(row):
    """Ngày đến hạn của kỳ kế tiếp, hoặc None nếu ngày bắt đầu không hợp lệ."""
    try:
        bd = datetime.strptime(row.get("ngay_bat_dau") or "", "%Y-%m-%d").date()
    except ValueError:
        return None
    return _add_months(bd, (row.get("da_tra") or 0) + 1)


def _norm_status(tt, so_ky, da_tra):
    """Chuẩn hoá trạng thái (dữ liệu cũ có thể là 'Dang tra' / 'Hoan thanh')."""
    t = (tt or "").strip().lower()
    if t in ("hoàn thành", "hoan thanh") or (so_ky > 0 and da_tra >= so_ky):
        return "Hoàn thành"
    return "Đang trả góp"


def _next_ma_tg(conn):
    n = (conn.execute("SELECT COALESCE(MAX(id),0) FROM tra_gop").fetchone()[0] or 0) + 1
    while conn.execute("SELECT 1 FROM tra_gop WHERE ma_tg=?", (f"TG{n:03d}",)).fetchone():
        n += 1
    return f"TG{n:03d}"


def _next_ma_tt(conn):
    n = (conn.execute("SELECT COALESCE(MAX(id),0) FROM thanh_toan").fetchone()[0] or 0) + 1
    while conn.execute("SELECT 1 FROM thanh_toan WHERE ma_tt=?", (f"TT{n:04d}",)).fetchone():
        n += 1
    return f"TT{n:04d}"


def _goc_da_tra(tong, lai_pct, so_ky, tien_ky, k):
    """Phần GỐC đã hoàn trả sau k kỳ (không tính phần lãi)."""
    if so_ky > 0 and k >= so_ky:
        return tong
    if k <= 0:
        return 0
    r = (lai_pct or 0) / 100
    if r <= 0:
        return tong * k / so_ky
    du_no = tong * (1 + r) ** k - tien_ky * ((1 + r) ** k - 1) / r
    return min(tong, max(0, tong - du_no))


def _dong_bo_don_hang(conn, tg_id):
    """Cập nhật đơn hàng gắn với hợp đồng: số tiền đã thanh toán + trạng thái thanh toán.
    (Màn Thanh toán đọc đúng 2 cột này của don_hang.) Không commit — người gọi commit."""
    tg = conn.execute("SELECT * FROM tra_gop WHERE id=?", (tg_id,)).fetchone()
    if not tg or not tg["don_hang_id"]:
        return
    dh = conn.execute("SELECT gia_ban_thuc FROM don_hang WHERE id=?", (tg["don_hang_id"],)).fetchone()
    if not dh:
        return
    gia = dh["gia_ban_thuc"] or 0
    so_ky = tg["so_thang"] or 0
    da = tg["so_thang_da_tra"] or 0
    truoc = tg["so_tien_tra_truoc"] or 0
    if so_ky > 0 and da >= so_ky:                      # trả hết góp → đơn đã thanh toán đủ
        da_tt, tt = gia, "Đã thanh toán"
    else:
        goc = _goc_da_tra(tg["tong_tien"] or 0, tg["lai_suat"], so_ky, tg["tien_hang_thang"] or 0, da)
        da_tt = min(gia, truoc + goc)
        tt = "Thanh toán một phần" if da_tt > 0 else "Chưa thanh toán"
    conn.execute("UPDATE don_hang SET so_tien_da_tt=?, trang_thai_tt=? WHERE id=?",
                 (int(round(da_tt)), tt, tg["don_hang_id"]))
    if tt == "Đã thanh toán":
        # Trạng thái xử lý của đơn (màn Đơn hàng): chỉ nâng "Chờ xử lý" → "Đã thanh toán".
        # Không đụng tới "Đã giao xe" / "Huỷ" / "Đặt cọc" để không hạ cấp đơn.
        conn.execute("UPDATE don_hang SET trang_thai='Đã thanh toán' "
                     "WHERE id=? AND trang_thai IN ('Chờ xử lý','Cho xu ly')", (tg["don_hang_id"],))


def _ghi_thanh_toan(conn, don_hang_id, so_tien, ghi_chu):
    """Thêm 1 dòng lịch sử vào bảng thanh_toan (tiền thực thu, gồm cả lãi)."""
    conn.execute(
        "INSERT INTO thanh_toan (ma_tt, don_hang_id, so_tien, phuong_thuc, trang_thai, ngay_tt, ghi_chu) "
        "VALUES (?,?,?,?,?,?,?)",
        (_next_ma_tt(conn), don_hang_id, so_tien, "Trả góp", "Hoàn thành",
         date.today().strftime("%Y-%m-%d"), ghi_chu))


STYLE = """
QWidget{font-family:'Segoe UI',Arial;}
QWidget#toolbar_widget{background:#ffffff;border-bottom:2px solid #e5e7eb;}
QPushButton#btn_add{background:qlineargradient(x1:0,y1:0,x2:1,y2:0,stop:0 #0891b2,stop:1 #06b6d4);
    color:white;border:none;border-radius:8px;font-size:13px;font-weight:800;padding:10px 20px;letter-spacing:0.5px;}
QPushButton#btn_add:hover{background:qlineargradient(x1:0,y1:0,x2:1,y2:0,stop:0 #06b6d4,stop:1 #5eead4);}

QPushButton#btn_pay{background:#dcfce7;color:#15803d;border:2px solid #86efac;border-radius:8px;padding:10px 18px;font-size:12px;font-weight:800;letter-spacing:0.5px;}
QPushButton#btn_pay:hover{background:#bbf7d0;border:2px solid #4ade80;}

QPushButton#btn_del{background:#fee2e2;color:#991b1b;border:2px solid #fca5a5;border-radius:8px;padding:10px 18px;font-size:12px;font-weight:800;letter-spacing:0.5px;}
QPushButton#btn_del:hover{background:#fecaca;border:2px solid #f87171;}
QLineEdit#search_box{background:#f9fafb;color:#111827;border:2px solid #e5e7eb;border-radius:8px;padding:9px 16px;font-size:13px;min-width:240px;font-weight:700;}
QLineEdit#search_box:focus{border:2px solid #2563eb;background:#ffffff;}
QTableWidget{background:#ffffff;alternate-background-color:#f9fafb;gridline-color:#f1f5f9;border:none;}
QTableWidget::item{padding:8px 12px;color:#1e293b;border-bottom:1px solid #f1f5f9;}
QTableWidget::item:selected{background:#eff6ff;color:#2563eb;}
QHeaderView::section{background:#f8fafc;color:#2563eb;font-size:14px;font-weight:900;letter-spacing:1px;padding:14px;border:none;border-bottom:2px solid #2563eb;}
QDialog{background:#ffffff;}
QLabel{color:#1e293b;background:transparent;font-size:14px;font-weight:600;}
QLineEdit,QComboBox,QDoubleSpinBox,QDateEdit{background:#f9fafb;color:#111827;border:1px solid #e5e7eb;border-radius:8px;padding:9px 12px;font-size:14px;}
QLineEdit:focus,QComboBox:focus,QDoubleSpinBox:focus,QDateEdit:focus{border-color:#2563eb;background:#ffffff;}
QPushButton#dlg_save{background:#2563eb;color:white;border:none;border-radius:9px;font-size:14px;font-weight:700;padding:11px 24px;}
QPushButton#dlg_save:hover{background:#1d4ed8;}
QPushButton#dlg_cancel{background:#f3f4f6;color:#6b7280;border:1px solid #d1d5db;border-radius:9px;font-size:13px;padding:10px 20px;}
QProgressBar{background:#f1f5f9;border-radius:6px;height:14px;text-align:center;color:#0284c7;font-size:11px;font-weight:800;border:1px solid #dbeafe;}
QProgressBar::chunk{background:qlineargradient(x1:0,y1:0,x2:1,y2:0,stop:0 #0891b2,stop:1 #06b6d4);border-radius:6px;}
"""


def _init_tragop_table():
    """Schema phải trùng với database.py (bảng tra_gop)."""
    conn = get_conn()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS tra_gop (
            id              INTEGER PRIMARY KEY AUTOINCREMENT,
            ma_tg           TEXT UNIQUE NOT NULL,
            don_hang_id     INTEGER REFERENCES don_hang(id),
            kh_id           INTEGER REFERENCES khach_hang(id),
            tong_tien       REAL NOT NULL DEFAULT 0,
            so_tien_tra_truoc REAL DEFAULT 0,
            lai_suat        REAL DEFAULT 0,
            so_thang        INTEGER DEFAULT 12,
            tien_hang_thang REAL DEFAULT 0,
            so_thang_da_tra INTEGER DEFAULT 0,
            tong_da_tra     REAL DEFAULT 0,
            con_lai         REAL DEFAULT 0,
            trang_thai      TEXT DEFAULT 'Dang tra',
            ngay_bat_dau    TEXT DEFAULT CURRENT_DATE,
            ghi_chu         TEXT,
            created_at      TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.commit(); conn.close()


class TraGopView(QWidget):
    def __init__(self, current_user=None):
        super().__init__()
        self.setObjectName("page_tra_gop")
        self.setStyleSheet(STYLE)
        self.current_user = current_user or {}
        self._rows = []; self._sel_id = None
        self._auto_checked = False
        _init_tragop_table()
        self._build()
        self._load()

    def _build(self):
        root = QVBoxLayout(self); root.setContentsMargins(0,0,0,0); root.setSpacing(0)

        # Toolbar
        tb = QWidget(); tb.setObjectName("toolbar_widget")
        tbh = QHBoxLayout(tb); tbh.setContentsMargins(16,10,16,10); tbh.setSpacing(8)
        title = QLabel("💳  Quản lý Thanh toán Trả góp")
        title.setStyleSheet("font-size:17px;font-weight:900;color:#111827;background:transparent;")
        btn_add = QPushButton("➕  Tạo hợp đồng TG"); btn_add.setObjectName("btn_add")
        btn_pay = QPushButton("💰  Ghi nhận trả kỳ"); btn_pay.setObjectName("btn_pay")
        btn_del = QPushButton("🗑  Xoá");
        btn_del.setObjectName("btn_del")
        btn_nhac = QPushButton("🔔  Nhắc nhở");
        btn_nhac.setObjectName("btn_pay")
        btn_excel = QPushButton("📊  Excel lịch trả");
        btn_excel.setObjectName("btn_add")
        btn_auto = QPushButton("⚡  Ghi nhận hàng loạt");
        btn_auto.setObjectName("btn_pay")
        self.search = QLineEdit(); self.search.setObjectName("search_box")
        self.search.setPlaceholderText("🔍  Tìm theo khách hàng, mã đơn...")
        self.search.textChanged.connect(lambda t: self._load(t.strip()))
        for b in [btn_add,btn_pay,btn_del,btn_nhac,btn_excel,btn_auto]:
            b.setCursor(Qt.CursorShape.PointingHandCursor)
        tbh.addWidget(title);
        tbh.addWidget(btn_add);
        tbh.addWidget(btn_pay)
        tbh.addWidget(btn_del);
        tbh.addWidget(btn_nhac); tbh.addWidget(btn_excel); tbh.addWidget(btn_auto)
        tbh.addStretch();
        tbh.addWidget(self.search)
        root.addWidget(tb)

        # Stat cards
        stat_w = QWidget();
        stat_w.setStyleSheet("background:#f9fafb;padding:14px 16px 12px;border-bottom:1px solid #e5e7eb;")
        stat_h = QHBoxLayout(stat_w); stat_h.setSpacing(10)
        self.sc_total  = self._stat("💳","Đang trả góp","0","#60a5fa")
        self.sc_hoan   = self._stat("✅","Hoàn thành","0","#4ade80")
        self.sc_dt_con = self._stat("💰","Tổng còn lại","0 ₫","#f87171")
        self.sc_dt_da  = self._stat("📊","Đã thu được","0 ₫","#a78bfa")
        for sc in [self.sc_total,self.sc_hoan,self.sc_dt_con,self.sc_dt_da]:
            stat_h.addWidget(sc)
        root.addWidget(stat_w)

        # Table
        cols = ["ID","ĐƠN HÀNG","KHÁCH HÀNG","TỔNG TIỀN","SỐ KỲ",
                "ĐÃ TRẢ","CÒN LẠI","TIẾN ĐỘ","LÃI SUẤT","TRẠNG THÁI"]
        self.tbl = QTableWidget(0, len(cols))
        self.tbl.setHorizontalHeaderLabels(cols)
        self.tbl.setAlternatingRowColors(True)
        self.tbl.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.tbl.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.tbl.setShowGrid(False); self.tbl.verticalHeader().setVisible(False)
        h = self.tbl.horizontalHeader()
        h.setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        h.hideSection(0)
        self.tbl.selectionModel().selectionChanged.connect(self._on_sel)
        root.addWidget(self.tbl)

        btn_add.clicked.connect(self._add)
        btn_pay.clicked.connect(self._pay)
        btn_del.clicked.connect(self._delete)
        btn_nhac.clicked.connect(self._nhac_nho)
        btn_excel.clicked.connect(self._export_lich)
        btn_auto.clicked.connect(self._ghi_nhan_hang_loat)

    def _stat(self, icon, label, val, color):
        w = QWidget()
        w.setStyleSheet(f"background:#ffffff;border-radius:12px;border:1px solid #e5e7eb;border-top:4px solid {color};")
        lv = QVBoxLayout(w); lv.setContentsMargins(14,12,14,12); lv.setSpacing(4)
        li = QLabel(icon); li.setStyleSheet("font-size:24px;background:transparent;")
        ll = QLabel(label); ll.setStyleSheet(f"font-size:13px;color:{color};font-weight:800;letter-spacing:.8px;background:transparent;text-transform:uppercase;")
        vl = QLabel(val); vl.setStyleSheet(f"font-size:20px;font-weight:900;color:{color};background:transparent;")
        lv.addWidget(li); lv.addWidget(ll); lv.addWidget(vl)
        w._val = vl; return w

    def _load(self, q=""):
        prev_sel = self._sel_id
        conn = get_conn()
        # Đặt alias để phần còn lại của view dùng tên cột quen thuộc (so_ky, da_tra, ...)
        sql = """SELECT tg.*,
                        tg.so_thang AS so_ky,
                        tg.tien_hang_thang AS tien_tra_moi_ky,
                        tg.so_thang_da_tra AS da_tra,
                        dh.ma_don, kh.ho_ten AS ten_kh
                 FROM tra_gop tg
                 LEFT JOIN don_hang dh ON tg.don_hang_id=dh.id
                 LEFT JOIN khach_hang kh ON tg.kh_id=kh.id"""
        p = []
        if q:
            sql += " WHERE kh.ho_ten LIKE ? OR dh.ma_don LIKE ? OR tg.ma_tg LIKE ?"
            p = [f"%{q}%"] * 3
        rows = [dict(r) for r in conn.execute(sql + " ORDER BY tg.id DESC", p).fetchall()]
        conn.close()

        today = date.today()
        for r in rows:
            r["so_ky"] = r["so_ky"] or 0
            r["da_tra"] = r["da_tra"] or 0
            r["tien_tra_moi_ky"] = r["tien_tra_moi_ky"] or 0
            r["tong_tien"] = r["tong_tien"] or 0
            r["lai_suat"] = r["lai_suat"] or 0
            r["con_lai_ky"] = max(0, r["so_ky"] - r["da_tra"])      # còn lại tính theo KỲ
            r["trang_thai"] = _norm_status(r.get("trang_thai"), r["so_ky"], r["da_tra"])
            r["trang_thai_hien_thi"] = r["trang_thai"]
            if r["trang_thai"] == "Đang trả góp":
                nd = _ngay_ky_tiep(r)
                if nd and nd < today:
                    r["trang_thai_hien_thi"] = "Trễ hạn"
        self._rows = rows

        # Stats
        total = sum(1 for r in rows if r["trang_thai"] == "Đang trả góp")
        hoan = sum(1 for r in rows if r["trang_thai"] == "Hoàn thành")
        dt_conlai = sum(r["con_lai_ky"] * r["tien_tra_moi_ky"] for r in rows if r["trang_thai"] == "Đang trả góp")
        dt_da = sum(r["da_tra"] * r["tien_tra_moi_ky"] for r in rows)
        self.sc_total._val.setText(str(total))
        self.sc_hoan._val.setText(str(hoan))
        self.sc_dt_con._val.setText(f"{dt_conlai/1e6:.0f} triệu")
        self.sc_dt_da._val.setText(f"{dt_da/1e6:.0f} triệu")

        STATUS_COL = {"Đang trả góp": "#60a5fa", "Hoàn thành": "#4ade80", "Trễ hạn": "#f87171"}
        self.tbl.selectionModel().blockSignals(True)   # tránh _on_sel chạy giữa lúc dựng lại bảng
        self.tbl.setRowCount(0)
        for row in rows:
            r = self.tbl.rowCount(); self.tbl.insertRow(r); self.tbl.setRowHeight(r, 50)
            so_ky = row["so_ky"] or 1
            da_tra = row["da_tra"]
            con_lai = row["con_lai_ky"]
            pct = min(100, int(da_tra / so_ky * 100))
            tien_ky = row["tien_tra_moi_ky"]

            vals = [str(row["id"]), row.get("ma_don") or "—",
                    row.get("ten_kh") or "—",
                    f"{row['tong_tien']/1e9:.3f} tỷ",
                    f"{so_ky} kỳ",
                    f"{da_tra}/{so_ky} kỳ  ({da_tra*tien_ky/1e6:.0f}tr)",
                    f"{con_lai} kỳ  ({con_lai*tien_ky/1e6:.0f}tr)",
                    f"{pct}%",
                    f"{row['lai_suat']}%/tháng",
                    row["trang_thai_hien_thi"]]

            for c, val in enumerate(vals):
                item = QTableWidgetItem(val)
                item.setData(Qt.ItemDataRole.UserRole, row["id"])
                item.setFont(QFont("Segoe UI", 13, QFont.Weight.Bold))
                if c == 7:
                    pb = QProgressBar()
                    pb.setValue(pct)
                    pb.setFormat(f"{pct}%")
                    self.tbl.setCellWidget(r, c, pb)
                    continue
                if c == 9:
                    item.setForeground(QColor(STATUS_COL.get(val, "#94a3b8")))
                if c in [3, 5, 6]:
                    color = "#16a34a" if c == 5 else "#dc2626" if c == 6 else "#1e293b"
                    item.setForeground(QColor(color))
                self.tbl.setItem(r, c, item)
        self.tbl.selectionModel().blockSignals(False)

        # Giữ lại dòng đang chọn nếu còn; nếu không thì bỏ chọn (tránh thao tác nhầm hợp đồng ẩn)
        self._sel_id = None
        if prev_sel is not None:
            for i, r in enumerate(rows):
                if r["id"] == prev_sel:
                    self.tbl.selectRow(i)
                    break

    def _on_sel(self, *args):
        items = self.tbl.selectedItems()
        self._sel_id = items[0].data(Qt.ItemDataRole.UserRole) if items else None

    def _selected_row(self):
        return next((r for r in self._rows if r["id"] == self._sel_id), None)

    def _add(self):
        if TraGopDialog(self).exec(): self._load()

    def _ghi_nhan_ky(self, conn, tg_id):
        """Ghi nhận thu 1 kỳ. Đọc lại số liệu MỚI từ DB (không dùng dữ liệu cũ trên màn hình).
        Trả về trạng thái mới, hoặc None nếu hợp đồng đã trả đủ / không tồn tại."""
        r = conn.execute(
            "SELECT ma_tg, don_hang_id, so_thang, so_thang_da_tra, tien_hang_thang FROM tra_gop WHERE id=?",
            (tg_id,)).fetchone()
        if not r: return None
        so, da, tien = r["so_thang"] or 0, r["so_thang_da_tra"] or 0, r["tien_hang_thang"] or 0
        if da >= so: return None
        da += 1
        tt = "Hoàn thành" if da >= so else "Đang trả góp"
        conn.execute(
            "UPDATE tra_gop SET so_thang_da_tra=?, tong_da_tra=?, con_lai=?, trang_thai=? WHERE id=?",
            (da, tien * da, tien * (so - da), tt, tg_id))
        if r["don_hang_id"]:
            _dong_bo_don_hang(conn, tg_id)
            _ghi_thanh_toan(conn, r["don_hang_id"], tien,
                            f"Trả góp {r['ma_tg']} — kỳ {da}/{so}")
        return tt

    def _pay(self):
        row = self._selected_row()
        if not row:
            QMessageBox.warning(self, "", "Chọn hợp đồng!"); return
        if row["trang_thai"] == "Hoàn thành" or row["con_lai_ky"] <= 0:
            QMessageBox.information(self, "", "Đã thanh toán hết!"); return
        if QMessageBox.question(self,
            "Ghi nhận thanh toán",
            f"Khách: {row.get('ten_kh') or ''}\n"
            f"Kỳ thứ: {row['da_tra']+1}/{row['so_ky']}\n"
            f"Số tiền: {row['tien_tra_moi_ky']/1e6:.1f} triệu ₫\n\nXác nhận đã nhận tiền?"
            + ("\n(Đơn hàng và màn Thanh toán sẽ được cập nhật theo)" if row.get("ma_don") else ""),
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        ) == QMessageBox.StandardButton.Yes:
            conn = get_conn()
            try:
                tt = self._ghi_nhan_ky(conn, row["id"])
                conn.commit()
            finally:
                conn.close()
            self._load()
            if tt == "Hoàn thành":
                msg = "Khách hàng đã trả hết góp!"
                if row.get("ma_don"):
                    msg += f"\nĐơn hàng {row['ma_don']} đã được ghi nhận thanh toán đủ."
                QMessageBox.information(self, "🎉 Hoàn thành!", msg)

    def _delete(self):
        row = self._selected_row()
        if not row:
            QMessageBox.warning(self, "", "Chọn hợp đồng!"); return
        msg = "Xoá hợp đồng trả góp này?"
        if row["trang_thai"] != "Hoàn thành" and row["da_tra"] > 0:
            msg = (f"Hợp đồng {row.get('ma_tg') or row['id']} đang trả dở "
                   f"({row['da_tra']}/{row['so_ky']} kỳ đã thu, còn {row['con_lai_ky']} kỳ).\n"
                   f"Xoá sẽ mất toàn bộ lịch sử thu tiền.\n\nVẫn xoá?")
        if QMessageBox.question(self, "Xác nhận", msg,
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No
        ) == QMessageBox.StandardButton.Yes:
            conn = get_conn(); conn.execute("DELETE FROM tra_gop WHERE id=?", (row["id"],))
            conn.commit(); conn.close(); self._sel_id = None; self._load()

    def _nhac_nho(self):
        """🔔 Nhắc nhở kỳ đến hạn trong 7 ngày tới"""
        from datetime import datetime, timedelta
        today = date.today()
        canh_bao = []
        qua_han = []

        for row in self._rows:
            if row.get("trang_thai") != "Đang trả góp": continue
            ngay_bd = row.get("ngay_bat_dau", "")
            if not ngay_bd: continue
            try:
                bd = datetime.strptime(ngay_bd, "%Y-%m-%d").date()
                da_tra = row.get("da_tra", 0) or 0
                # Ngày kỳ tiếp theo = ngày bắt đầu + (số kỳ đã trả + 1) tháng
                ngay_ky_tiep = _add_months(bd, da_tra + 1)
                delta = (ngay_ky_tiep - today).days
                ten_kh = row.get("ten_kh", "") or "—"
                tien_ky = row.get("tien_tra_moi_ky", 0) or 0
                ky_so = da_tra + 1
                so_ky = row.get("so_ky", 0)
                info = (f"👤 {ten_kh}  |  Kỳ {ky_so}/{so_ky}  "
                        f"|  {tien_ky / 1e6:.1f} triệu  |  "
                        f"Ngày: {ngay_ky_tiep.strftime('%d/%m/%Y')}")
                if delta < 0:
                    qua_han.append(f"❌ Quá hạn {abs(delta)} ngày — {info}")
                elif delta <= 7:
                    canh_bao.append(f"⚠️ Còn {delta} ngày — {info}")
            except Exception:
                continue

        # Hiển thị dialog nhắc nhở
        dlg = QDialog(self)
        dlg.setWindowTitle("🔔 Nhắc nhở kỳ trả góp")
        dlg.setMinimumWidth(620)
        dlg.setStyleSheet("""
            QDialog { background: #ffffff; }
            QLabel { background: transparent; }
            QPushButton { background: #2563eb; color: white; border: none;
                border-radius: 8px; padding: 10px 24px; font-size: 13px; font-weight: 700; }
            QPushButton:hover { background: #1d4ed8; }
        """)
        lv = QVBoxLayout(dlg)
        lv.setContentsMargins(24, 20, 24, 20)
        lv.setSpacing(12)

        title = QLabel("🔔  Nhắc nhở kỳ trả góp đến hạn")
        title.setStyleSheet("font-size:16px;font-weight:800;color:#0f172a;")
        lv.addWidget(title)

        sep = QFrame();
        sep.setFrameShape(QFrame.Shape.HLine)
        sep.setStyleSheet("background:#e5e7eb;max-height:1px;")
        lv.addWidget(sep)

        if not qua_han and not canh_bao:
            ok_lbl = QLabel("✅  Không có kỳ nào đến hạn trong 7 ngày tới!")
            ok_lbl.setStyleSheet("color:#059669;font-size:14px;font-weight:700;")
            lv.addWidget(ok_lbl)
        else:
            if qua_han:
                lbl_qh = QLabel(f"❌  QUÁ HẠN — {len(qua_han)} hợp đồng")
                lbl_qh.setStyleSheet("color:#dc2626;font-size:13px;font-weight:800;")
                lv.addWidget(lbl_qh)
                for txt in qua_han:
                    row_w = QLabel(txt)
                    row_w.setStyleSheet(
                        "color:#991b1b;font-size:12px;font-weight:600;"
                        "background:#fff1f2;border-radius:6px;padding:8px 12px;"
                        "border-left:4px solid #dc2626;")
                    row_w.setWordWrap(True)
                    lv.addWidget(row_w)

            if canh_bao:
                lbl_cb = QLabel(f"⚠️  SẮP ĐẾN HẠN (7 ngày) — {len(canh_bao)} hợp đồng")
                lbl_cb.setStyleSheet("color:#d97706;font-size:13px;font-weight:800;margin-top:8px;")
                lv.addWidget(lbl_cb)
                for txt in canh_bao:
                    row_w = QLabel(txt)
                    row_w.setStyleSheet(
                        "color:#92400e;font-size:12px;font-weight:600;"
                        "background:#fffbeb;border-radius:6px;padding:8px 12px;"
                        "border-left:4px solid #f59e0b;")
                    row_w.setWordWrap(True)
                    lv.addWidget(row_w)

        lv.addSpacing(8)
        btn_ok = QPushButton("✓  Đã xem")
        btn_ok.clicked.connect(dlg.accept)
        bh = QHBoxLayout();
        bh.addStretch();
        bh.addWidget(btn_ok)
        lv.addLayout(bh)
        dlg.exec()

    def _export_lich(self):
        """📊 Xuất Excel lịch trả kỳ chi tiết"""
        if not self._sel_id:
            QMessageBox.warning(self, "", "Chọn hợp đồng cần xuất lịch!");
            return

        row = next((r for r in self._rows if r["id"] == self._sel_id), None)
        if not row: return

        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
            import datetime as dt

            ten_kh = row.get("ten_kh", "KH") or "KH"
            tong = row.get("tong_tien", 0) or 0
            so_ky = row.get("so_ky", 12) or 12
            tien_ky = row.get("tien_tra_moi_ky", 0) or 0
            da_tra = row.get("da_tra", 0) or 0
            lai = row.get("lai_suat", 0) or 0
            ngay_bd_str = row.get("ngay_bat_dau") or dt.date.today().strftime("%Y-%m-%d")
            try:
                ngay_bd = dt.datetime.strptime(ngay_bd_str, "%Y-%m-%d").date()
            except Exception:
                ngay_bd = dt.date.today()

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Lich tra gop"

            thin = Side(style="thin", color="CBD5E1")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            # ── Tiêu đề ──────────────────────────────────────
            ws.merge_cells("A1:G1")
            ws["A1"] = f"LỊCH TRẢ GÓP — {ten_kh.upper()}"
            ws["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
            ws["A1"].fill = PatternFill("solid", fgColor="0891B2")
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 36

            # ── Thông tin hợp đồng ───────────────────────────
            info = [
                ("Khách hàng", ten_kh),
                ("Tổng vay", f"{tong / 1e9:.3f} tỷ ₫"),
                ("Số kỳ", f"{so_ky} kỳ"),
                ("Tiền mỗi kỳ", f"{tien_ky / 1e6:.2f} triệu ₫"),
                ("Lãi suất", f"{lai}%/tháng"),
                ("Ngày bắt đầu", ngay_bd_str),
            ]
            for i, (k, v) in enumerate(info, 2):
                ws.merge_cells(f"A{i}:C{i}")
                ws[f"A{i}"] = k
                ws[f"A{i}"].font = Font(name="Segoe UI", size=10, bold=True, color="0891B2")
                ws[f"A{i}"].alignment = Alignment(horizontal="right", vertical="center")
                ws.merge_cells(f"D{i}:G{i}")
                ws[f"D{i}"] = v
                ws[f"D{i}"].font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
                ws.row_dimensions[i].height = 20

            # ── Header bảng ──────────────────────────────────
            hr = len(info) + 2
            headers = ["KỲ", "NGÀY ĐẾN HẠN", "SỐ TIỀN", "GỐC", "LÃI",
                       "TRẠNG THÁI", "GHI CHÚ"]
            col_widths = [6, 16, 16, 16, 14, 14, 20]
            for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
                cell = ws.cell(row=hr, column=ci, value=h)
                cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="0891B2")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border
                ws.column_dimensions[get_column_letter(ci)].width = w
            ws.row_dimensions[hr].height = 28

            # ── Dữ liệu từng kỳ ─────────────────────────────
            fill_paid = PatternFill("solid", fgColor="F0FDF4")
            fill_due = PatternFill("solid", fgColor="FFFBEB")
            fill_overdue = PatternFill("solid", fgColor="FFF1F2")
            fill_normal = PatternFill("solid", fgColor="FFFFFF")
            fill_alt = PatternFill("solid", fgColor="F0F9FF")
            today = dt.date.today()

            # Tính gốc và lãi mỗi kỳ
            du_no = tong
            lai_thang = lai / 100

            for ky in range(1, so_ky + 1):
                ri = hr + ky
                ws.row_dimensions[ri].height = 22

                ngay_ky = _add_months(ngay_bd, ky)
                tien_lai = du_no * lai_thang
                tien_goc = tien_ky - tien_lai if lai_thang > 0 else tien_ky
                du_no = max(0, du_no - tien_goc)

                delta = (ngay_ky - today).days

                if ky <= da_tra:
                    tt = "✅ Đã trả"
                    fill = fill_paid
                    tt_color = "059669"
                elif delta < 0:
                    tt = "❌ Quá hạn"
                    fill = fill_overdue
                    tt_color = "DC2626"
                elif delta <= 7:
                    tt = "⚠️ Sắp đến"
                    fill = fill_due
                    tt_color = "D97706"
                else:
                    tt = "⏳ Chờ"
                    fill = fill_normal if ky % 2 == 0 else fill_alt
                    tt_color = "64748B"

                vals = [
                    ky,
                    ngay_ky.strftime("%d/%m/%Y"),
                    f"{tien_ky / 1e6:.2f} tr",
                    f"{tien_goc / 1e6:.2f} tr",
                    f"{tien_lai / 1e6:.2f} tr",
                    tt,
                    "Kỳ hiện tại" if ky == da_tra + 1 else "",
                ]

                for ci, val in enumerate(vals, 1):
                    cell = ws.cell(row=ri, column=ci, value=val)
                    cell.fill = fill
                    cell.border = border
                    cell.alignment = Alignment(vertical="center",
                                               horizontal="center" if ci in (1, 2, 3, 4, 5, 6) else "left")
                    if ci == 1:
                        cell.font = Font(name="Segoe UI", size=11,
                                         bold=(ky == da_tra + 1), color="0891B2")
                    elif ci == 6:
                        cell.font = Font(name="Segoe UI", size=11,
                                         bold=True, color=tt_color)
                    elif ci == 7 and ky == da_tra + 1:
                        cell.font = Font(name="Segoe UI", size=10,
                                         bold=True, color="D97706")
                    else:
                        cell.font = Font(name="Segoe UI", size=11, color="374151")

            ws.freeze_panes = f"A{hr + 1}"

            import re
            ten_safe = re.sub(r'[^\w]', '_', ten_kh)
            default_name = f"LichTraGop_{ten_safe}_{dt.datetime.now().strftime('%d%m%Y_%H%M')}.xlsx"
            fname, _ = QFileDialog.getSaveFileName(self, "Lưu lịch trả góp", default_name,
                                                   "Excel (*.xlsx)")
            if not fname:
                return
            wb.save(fname)
            QMessageBox.information(self, "✅ Xuất thành công!",
                                    f"Đã xuất lịch trả góp:\n{fname}\n\n"
                                    f"📋 {so_ky} kỳ  |  ✅ Đã trả: {da_tra}  |  ⏳ Còn lại: {so_ky - da_tra}")

        except ImportError:
            QMessageBox.critical(self, "Lỗi",
                                 "Cần cài thêm:\npip install openpyxl")
        except Exception as e:
            import traceback
            QMessageBox.critical(self, "Lỗi", f"{str(e)}\n\n{traceback.format_exc()}")

    def showEvent(self, e):
        super().showEvent(e)
        if not self._auto_checked:          # chỉ kiểm tra 1 lần, lúc mở tab lần đầu
            self._auto_checked = True
            QTimer.singleShot(300, self._auto_check_khi_mo_app)

    def _auto_check_khi_mo_app(self):
        """Báo các kỳ đã đến hạn / quá hạn khi mở tab. KHÔNG tự ghi nhận đã thu tiền:
        chỉ mở danh sách để người dùng tick xác nhận từng hợp đồng."""
        today = date.today()
        den_han = []
        for row in self._rows:
            if row["trang_thai"] != "Đang trả góp": continue
            nd = _ngay_ky_tiep(row)
            if nd and (nd - today).days <= 0:
                den_han.append(row)
        if not den_han:
            return
        ds = "\n".join(
            f"• {r.get('ten_kh') or '—'}  —  Kỳ {r['da_tra'] + 1}/{r['so_ky']}  —  {r['tien_tra_moi_ky'] / 1e6:.1f} triệu"
            for r in den_han[:10])
        if len(den_han) > 10:
            ds += f"\n… và {len(den_han) - 10} hợp đồng khác"
        ret = QMessageBox.question(
            self, "🔔 Kỳ trả góp đến hạn!",
            f"Có {len(den_han)} hợp đồng đã đến hạn / quá hạn:\n\n{ds}\n\n"
            f"Mở danh sách ghi nhận hàng loạt để xác nhận thu tiền?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
        if ret == QMessageBox.StandardButton.Yes:
            self._ghi_nhan_hang_loat()

    def _ghi_nhan_hang_loat(self):
        """Ghi nhận nhiều kỳ trả cùng lúc"""
        from datetime import date
        today = date.today()

        # Dialog chọn hợp đồng
        dlg = QDialog(self)
        dlg.setWindowTitle("⚡ Ghi nhận hàng loạt")
        dlg.setMinimumWidth(680)
        dlg.setStyleSheet("""
            QDialog { background:#ffffff; }
            QLabel { background:transparent; color:#0f172a; }
            QTableWidget { background:#ffffff; border:1px solid #e5e7eb;
                alternate-background-color:#f8fafc; gridline-color:#f1f5f9; }
            QTableWidget::item { padding:8px; color:#1e293b; }
            QTableWidget::item:selected { background:#eff6ff; color:#2563eb; }
            QHeaderView::section { background:#f8fafc; color:#2563eb;
                font-weight:800; padding:10px; border:none;
                border-bottom:2px solid #2563eb; }
            QPushButton#btn_ok { background:#2563eb; color:white; border:none;
                border-radius:8px; padding:10px 24px; font-size:13px; font-weight:700; }
            QPushButton#btn_ok:hover { background:#1d4ed8; }
            QPushButton#btn_cancel { background:#f3f4f6; color:#6b7280;
                border:1px solid #d1d5db; border-radius:8px;
                padding:10px 20px; font-size:13px; }
        """)

        lv = QVBoxLayout(dlg)
        lv.setContentsMargins(20, 20, 20, 20)
        lv.setSpacing(12)

        title = QLabel("⚡  Chọn hợp đồng cần ghi nhận trả kỳ")
        title.setStyleSheet("font-size:15px;font-weight:800;color:#0f172a;")
        lv.addWidget(title)

        note = QLabel("✅ Tick chọn các hợp đồng đã nhận tiền — bấm Xác nhận để ghi nhận tất cả")
        note.setStyleSheet("font-size:12px;color:#64748b;font-weight:600;")
        lv.addWidget(note)

        from PyQt6.QtWidgets import QTableWidget, QTableWidgetItem, QHeaderView
        from PyQt6.QtCore import Qt

        tbl = QTableWidget(0, 5)
        tbl.setHorizontalHeaderLabels(["☑", "KHÁCH HÀNG", "KỲ", "SỐ TIỀN", "HẠN TRẢ"])
        tbl.setAlternatingRowColors(True)
        tbl.setShowGrid(False)
        tbl.verticalHeader().setVisible(False)
        tbl.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        tbl.setMinimumHeight(250)

        dang_tra = [r for r in self._rows if r.get("trang_thai") == "Đang trả góp"]

        for row in dang_tra:
            r = tbl.rowCount();
            tbl.insertRow(r);
            tbl.setRowHeight(r, 44)

            # Tính ngày kỳ tiếp
            try:
                from datetime import datetime as dt2
                bd = dt2.strptime(row.get("ngay_bat_dau", "2026-01-01"), "%Y-%m-%d").date()
                da_tra = row.get("da_tra", 0) or 0
                ngay_ky = _add_months(bd, da_tra + 1)
                delta = (ngay_ky - today).days
                ngay_str = ngay_ky.strftime("%d/%m/%Y")
                if delta < 0:
                    ngay_str += f" (quá {abs(delta)}ngày)"
                elif delta == 0:
                    ngay_str += " (HÔM NAY)"
                elif delta <= 7:
                    ngay_str += f" (còn {delta}ngày)"
            except Exception:
                ngay_str = "—"
                delta = 99

            # Checkbox
            chk = QTableWidgetItem()
            chk.setCheckState(Qt.CheckState.Checked if delta <= 0
                              else Qt.CheckState.Unchecked)
            chk.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            tbl.setItem(r, 0, chk)

            # Tên KH
            i1 = QTableWidgetItem(row.get("ten_kh", "—") or "—")
            i1.setFont(QFont("Segoe UI", 12, QFont.Weight.Bold))
            i1.setForeground(QColor("#0f172a"))
            tbl.setItem(r, 1, i1)

            # Kỳ
            ky_so = (row.get("da_tra") or 0) + 1
            i2 = QTableWidgetItem(f"Kỳ {ky_so}/{row.get('so_ky', 0)}")
            i2.setFont(QFont("Segoe UI", 11))
            i2.setForeground(QColor("#7c3aed"))
            i2.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            tbl.setItem(r, 2, i2)

            # Số tiền
            tien = row.get("tien_tra_moi_ky", 0) or 0
            i3 = QTableWidgetItem(f"{tien / 1e6:.1f} triệu")
            i3.setFont(QFont("Segoe UI", 11, QFont.Weight.Bold))
            i3.setForeground(QColor("#059669"))
            i3.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            tbl.setItem(r, 3, i3)

            # Hạn trả
            i4 = QTableWidgetItem(ngay_str)
            i4.setFont(QFont("Segoe UI", 11))
            color = "#dc2626" if delta < 0 else "#d97706" if delta <= 7 else "#64748b"
            i4.setForeground(QColor(color))
            i4.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
            tbl.setItem(r, 4, i4)
            tbl.setItem(r, 0, chk)

        lv.addWidget(tbl)

        # Footer
        bh = QHBoxLayout();
        bh.addStretch()
        btn_cancel = QPushButton("Huỷ");
        btn_cancel.setObjectName("btn_cancel")
        btn_cancel.clicked.connect(dlg.reject)
        btn_ok = QPushButton("⚡  Xác nhận ghi nhận");
        btn_ok.setObjectName("btn_ok")

        def _do_confirm():
            selected = []
            for r in range(tbl.rowCount()):
                chk_item = tbl.item(r, 0)
                if chk_item and chk_item.checkState() == Qt.CheckState.Checked:
                    selected.append(dang_tra[r])
            if not selected:
                QMessageBox.warning(dlg, "", "Chưa chọn hợp đồng nào!");
                return

            ten_list = "\n".join([f"• {r.get('ten_kh', '—')}" for r in selected])
            if QMessageBox.question(dlg, "Xác nhận",
                                    f"Ghi nhận đã thu tiền {len(selected)} hợp đồng:\n\n{ten_list}",
                                    QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
                                    ) == QMessageBox.StandardButton.Yes:
                conn = get_conn()
                hoan_thanh = 0
                try:
                    for row in selected:
                        if self._ghi_nhan_ky(conn, row["id"]) == "Hoàn thành":
                            hoan_thanh += 1
                    conn.commit()
                finally:
                    conn.close()
                self._load()
                dlg.accept()
                msg = f"✅ Đã ghi nhận {len(selected)} kỳ trả góp!"
                if hoan_thanh:
                    msg += f"\n🎉 {hoan_thanh} hợp đồng hoàn thành!"
                QMessageBox.information(self, "✅ Thành công!", msg)

        btn_ok.clicked.connect(_do_confirm)
        bh.addWidget(btn_cancel);
        bh.addWidget(btn_ok)
        lv.addLayout(bh)
        dlg.exec()



    def refresh(self):
        self._load()


class TraGopDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Tạo hợp đồng trả góp")
        self.setMinimumWidth(500)
        self.setStyleSheet(STYLE)
        self._moi_ky = 0
        self._dh_info = {}          # don_hang_id -> (kh_id, gia_ban_thuc)
        self._build()

    def _build(self):
        outer = QVBoxLayout(self); outer.setContentsMargins(24,20,24,20); outer.setSpacing(14)
        title = QLabel("💳  HỢP ĐỒNG TRẢ GÓP")
        title.setStyleSheet("font-size:16px;font-weight:900;color:#0284c7;background:transparent;")
        sep = QFrame();
        sep.setFrameShape(QFrame.Shape.HLine)
        sep.setStyleSheet("background:#e5e7eb;max-height:2px;border-radius:1px;")
        outer.addWidget(title); outer.addWidget(sep)

        form = QFormLayout(); form.setSpacing(10); form.setLabelAlignment(Qt.AlignmentFlag.AlignRight)

        conn = get_conn()
        # Chỉ liệt kê đơn trả góp CHƯA có hợp đồng
        dh_rows = conn.execute("""
            SELECT dh.id, dh.ma_don, dh.kh_id, kh.ho_ten, dh.gia_ban_thuc
            FROM don_hang dh JOIN khach_hang kh ON dh.kh_id=kh.id
            WHERE dh.phuong_thuc IN ('Trả góp','Tra gop')
              AND dh.id NOT IN (SELECT don_hang_id FROM tra_gop WHERE don_hang_id IS NOT NULL)
            ORDER BY dh.id DESC""").fetchall()
        kh_rows = conn.execute("SELECT id,ma_kh,ho_ten FROM khach_hang ORDER BY id DESC").fetchall()
        conn.close()

        self.f_dh = QComboBox()
        self.f_dh.addItem("— Không gắn đơn hàng —", None)
        for r in dh_rows:
            gia = r["gia_ban_thuc"] or 0
            self._dh_info[r["id"]] = (r["kh_id"], gia)
            self.f_dh.addItem(f"{r['ma_don']} — {r['ho_ten']} ({gia/1e9:.2f} tỷ)", r["id"])
        self.f_kh = QComboBox()
        for r in kh_rows: self.f_kh.addItem(f"{r[1]} — {r[2]}", r[0])
        self.f_truoc = QDoubleSpinBox()
        self.f_truoc.setRange(0,10e9); self.f_truoc.setSingleStep(50e6)
        self.f_truoc.setDecimals(0); self.f_truoc.setSuffix(" ₫")
        self.f_tong = QDoubleSpinBox()
        self.f_tong.setRange(0,10e9); self.f_tong.setSingleStep(50e6)
        self.f_tong.setDecimals(0); self.f_tong.setSuffix(" ₫")
        self.f_ky = QComboBox()
        self.f_ky.addItems(["6","12","18","24","36","48","60"])
        self.f_ky.setCurrentText("12")
        self.f_lai = QDoubleSpinBox()
        self.f_lai.setRange(0,5); self.f_lai.setSingleStep(0.1)
        self.f_lai.setValue(0.8); self.f_lai.setSuffix("%/tháng")
        self.f_ngay = QDateEdit(QDate.currentDate())
        self.f_ngay.setDisplayFormat("yyyy-MM-dd")
        self.f_ngay.setCalendarPopup(True)
        self.f_ghi = QLineEdit(); self.f_ghi.setPlaceholderText("Ghi chú...")

        def lbl(t, color="#0284c7"):
            l = QLabel(t)
            l.setStyleSheet(f"color:{color};font-size:12px;font-weight:800;background:transparent;letter-spacing:.8px;")
            return l

        for (txt, color), w in zip([("Đơn hàng", "#0284c7"), ("Khách hàng *", "#0284c7"),
                                    ("Trả trước", "#0284c7"), ("Tổng tiền vay *", "#dc2626"),
                                    ("Số kỳ thanh toán", "#f97316"),
                                    ("Lãi suất", "#059669"), ("Ngày bắt đầu", "#7c3aed"),
                                    ("Ghi chú", "#6b7280")],
                                   [self.f_dh, self.f_kh, self.f_truoc, self.f_tong, self.f_ky,
                                    self.f_lai, self.f_ngay, self.f_ghi]):
            form.addRow(lbl(txt, color), w)

        # Auto tính tiền mỗi kỳ
        self.lbl_tieng = QLabel("💰 Mỗi kỳ: 0 ₫")
        self.lbl_tieng.setStyleSheet("color:#4ade80;font-size:13px;font-weight:700;background:#052e16;border-radius:8px;padding:8px 12px;")
        form.addRow("", self.lbl_tieng)
        self.f_dh.currentIndexChanged.connect(lambda _=None: self._on_dh_changed())
        self.f_truoc.valueChanged.connect(lambda _=None: self._on_dh_changed())
        self.f_tong.valueChanged.connect(self._calc)
        self.f_ky.currentTextChanged.connect(self._calc)
        self.f_lai.valueChanged.connect(self._calc)
        outer.addLayout(form)

        bh = QHBoxLayout(); bh.addStretch()
        bc = QPushButton("Huỷ bỏ"); bc.setObjectName("dlg_cancel"); bc.clicked.connect(self.reject)
        bs = QPushButton("💳  Tạo hợp đồng"); bs.setObjectName("dlg_save")
        bs.clicked.connect(self._save); bs.setDefault(True)
        bh.addWidget(bc); bh.addWidget(bs); outer.addLayout(bh)
        self._on_dh_changed()
        self._calc()

    def _on_dh_changed(self):
        """Chọn đơn hàng → tự điền khách hàng và tổng vay (= giá bán − trả trước)."""
        info = self._dh_info.get(self.f_dh.currentData())
        if not info:
            return
        kh_id, gia = info
        idx = self.f_kh.findData(kh_id)
        if idx >= 0:
            self.f_kh.setCurrentIndex(idx)
        self.f_tong.setValue(max(0, gia - self.f_truoc.value()))

    def _calc(self):
        try:
            tong = self.f_tong.value()
            ky = int(self.f_ky.currentText())
            lai = self.f_lai.value()

            moi_ky = tinh_tien_tra_gop(tong, lai, ky)
            self.lbl_tieng.setText(
                f"💰 Mỗi kỳ: {moi_ky/1e6:.2f} triệu ₫  |  "
            f"Tổng: {moi_ky*ky/1e6:.1f} triệu"
            )
            self._moi_ky = moi_ky
        except Exception:
            self._moi_ky = 0

    def _save(self):
        kh_id = self.f_kh.currentData()
        tong = self.f_tong.value()
        ky = int(self.f_ky.currentText())
        if not kh_id or tong <= 0:
            QMessageBox.warning(self, "", "Điền đủ thông tin!"); return
        self._calc()
        moi_ky = self._moi_ky
        if moi_ky <= 0:
            QMessageBox.warning(self, "", "Không tính được tiền mỗi kỳ — kiểm tra lại số tiền vay / số kỳ / lãi suất!"); return
        dh_id = self.f_dh.currentData()
        conn = get_conn()
        try:
            if dh_id is not None and conn.execute(
                    "SELECT 1 FROM tra_gop WHERE don_hang_id=?", (dh_id,)).fetchone():
                QMessageBox.warning(self, "", "Đơn hàng này đã có hợp đồng trả góp!"); return
            ma_tg = _next_ma_tg(conn)
            cur = conn.execute("""INSERT INTO tra_gop
                (ma_tg, don_hang_id, kh_id, tong_tien, so_tien_tra_truoc, lai_suat, so_thang,
                 tien_hang_thang, so_thang_da_tra, tong_da_tra, con_lai, trang_thai,
                 ngay_bat_dau, ghi_chu)
                VALUES (?,?,?,?,?,?,?,?,0,0,?,?,?,?)""",
                (ma_tg, dh_id, kh_id, tong, self.f_truoc.value(), self.f_lai.value(), ky,
                 moi_ky, moi_ky * ky, "Đang trả góp",
                 self.f_ngay.date().toString("yyyy-MM-dd"), self.f_ghi.text().strip()))
            if dh_id is not None:
                _dong_bo_don_hang(conn, cur.lastrowid)          # ghi nhận tiền trả trước (nếu có)
                if self.f_truoc.value() > 0:
                    _ghi_thanh_toan(conn, dh_id, self.f_truoc.value(), f"Trả trước hợp đồng {ma_tg}")
            conn.commit()
            QMessageBox.information(self, "✅ OK", f"Tạo hợp đồng trả góp {ma_tg} thành công!")
            self.accept()
        except Exception as e: QMessageBox.critical(self, "Lỗi", str(e))
        finally: conn.close()