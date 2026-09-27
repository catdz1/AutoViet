"""
views/nhan_vien_view.py — Màn hình quản lý Nhân viên (+ hồ sơ, AI chấm công)
Tách từ views/other_views.py để mỗi màn hình nghiệp vụ nằm 1 file riêng.
"""
import os                  # ← THÊM DÒNG NÀY
import pandas as pd
import calendar, random
from datetime import datetime, date
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QTableWidget, QTableWidgetItem,
    QPushButton, QLineEdit, QDialog, QFormLayout, QComboBox,
    QMessageBox, QHeaderView, QLabel, QDoubleSpinBox, QTextEdit,
    QFrame, QApplication, QTabWidget, QScrollArea, QGridLayout, QInputDialog
)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QColor, QFont, QPixmap   # ← thêm QPixmap vào đây luôn
from database import get_conn

from views.base_widgets import BaseDialog, BaseView

# ════════════════════════════════════════════════════════════════
# NHÂN VIÊN
# ════════════════════════════════════════════════════════════════
class NhanVienView(BaseView):
    COLS=[("MÃ NV","ma_nv",False),("HỌ TÊN","ho_ten",True),("CHỨC VỤ","chuc_vu",False),
          ("SỐ ĐT","so_dt",False),("LƯƠNG","luong",False),("TRẠNG THÁI","trang_thai",False)]
    def __init__(self,current_user=None): super().__init__("page_nhan_vien",current_user)

    def _build_toolbar(self):
        self._btn("+ Thêm NV", "btn_add", self._add);
        self._btn("✏  Sửa", None, self._edit)
        self._btn("🗑  Xoá", "btn_del", self._delete);
        self._btn("📊 Excel", "btn_excel", self._export)
        self._btn("👁 Xem hồ sơ", None, self._xem_ho_so)  # ← thêm
        self._btn("🤖 AI Chấm công", "btn_add", self._ai_cham_cong)  # ← thêm
        self._search_box("🔍 Tìm nhân viên...")
    def _load(self,q=""):
        conn=get_conn(); sql="SELECT * FROM nhan_vien"; p=[]
        if q: sql+=" WHERE ho_ten LIKE ? OR ma_nv LIKE ? OR chuc_vu LIKE ?"; p=[f"%{q}%"]*3
        self._rows=[dict(r) for r in conn.execute(sql+" ORDER BY id DESC",p).fetchall()]
        conn.close(); self._render()
        STATUS={"Đang làm":"#4ade80","Thử việc":"#fbbf24","Nghỉ việc":"#f87171"}
        for r,row in enumerate(self._rows):
            i=QTableWidgetItem(f"{int(row.get('luong',0) or 0):,} ₫")
            i.setForeground(QColor("#4ade80")); i.setFont(QFont("Segoe UI",12,QFont.Weight.Bold))
            self.tbl.setItem(r,4,i)
            ti=QTableWidgetItem(row.get("trang_thai",""))
            ti.setForeground(QColor(STATUS.get(row.get("trang_thai",""),"#94a3b8")))
            self.tbl.setItem(r,5,ti)
    def _add(self):
        if NhanVienDialog(self).exec(): self._load()
    def _edit(self):
        if not self._check_sel("sửa"): return
        if NhanVienDialog(self,self._get_row("nhan_vien",self._sel_id)).exec(): self._load()
    def _delete(self):
        if not self._check_sel("xoá"): return
        conn = get_conn()
        so_don = conn.execute("SELECT COUNT(*) FROM don_hang WHERE nv_id=?",(self._sel_id,)).fetchone()[0]
        so_dv  = conn.execute("SELECT COUNT(*) FROM dich_vu WHERE nv_id=?",(self._sel_id,)).fetchone()[0]
        so_cc  = conn.execute("SELECT COUNT(*) FROM cham_cong WHERE nv_id=?",(self._sel_id,)).fetchone()[0]
        conn.close()
        if so_don > 0:
            QMessageBox.critical(self,"❌ Không thể xoá",
                f"Nhân viên đang có {so_don} đơn hàng!\nVui lòng chuyển đơn hàng sang NV khác trước.")
            return
        if so_dv > 0:
            QMessageBox.critical(self,"❌ Không thể xoá",
                f"Nhân viên đang có {so_dv} phiếu dịch vụ!\nVui lòng chuyển phiếu sang NV khác trước.")
            return
        if QMessageBox.question(self,"Xác nhận xoá",
            f"Xoá nhân viên này?\n\n⚠️ Dữ liệu chấm công ({so_cc} bản ghi) cũng sẽ bị xoá!",
            QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No
        )==QMessageBox.StandardButton.Yes:
            try:
                conn=get_conn()
                conn.execute("DELETE FROM cham_cong WHERE nv_id=?",(self._sel_id,))
                conn.execute("DELETE FROM nhan_vien WHERE id=?",(self._sel_id,))
                conn.commit(); conn.close()
                self._sel_id=None; self._load()
                QMessageBox.information(self,"✅ Thành công","Đã xoá nhân viên!")
            except Exception as e:
                QMessageBox.critical(self,"❌ Lỗi",f"Không thể xoá:\n{str(e)}")

    def _export(self):
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        import datetime

        conn = get_conn()
        rows = conn.execute("SELECT ma_nv,ho_ten,chuc_vu,so_dt,email,luong,trang_thai FROM nhan_vien").fetchall()
        conn.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Nhân viên"

        ws.merge_cells("A1:G1")
        ws["A1"] = "DANH SÁCH NHÂN VIÊN — AUTOVIET"
        ws["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="1E40AF")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        ws.merge_cells("A2:G2")
        ws["A2"] = f"Ngày xuất: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
        ws["A2"].alignment = Alignment(horizontal="right")
        ws.row_dimensions[2].height = 20

        headers = ["MÃ NV", "HỌ TÊN", "CHỨC VỤ", "SỐ ĐT", "EMAIL", "LƯƠNG", "TRẠNG THÁI"]
        col_widths = [12, 25, 18, 16, 28, 16, 14]
        thin = Side(style="thin", color="BFDBFE")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for i, (h, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=3, column=i, value=h)
            cell.font = Font(name="Segoe UI", size=11, bold=True, color="1E40AF")
            cell.fill = PatternFill("solid", fgColor="DBEAFE")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[3].height = 28

        fill_white = PatternFill("solid", fgColor="FFFFFF")
        fill_alt   = PatternFill("solid", fgColor="F0F7FF")
        font_data  = Font(name="Segoe UI", size=11, color="1E293B")
        font_ma    = Font(name="Segoe UI", size=11, bold=True, color="2563EB")

        for ri, row in enumerate(rows, 4):
            fill = fill_white if ri % 2 == 0 else fill_alt
            ws.row_dimensions[ri].height = 24
            for ci, val in enumerate(row, 1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.fill = fill
                cell.border = border
                cell.alignment = Alignment(vertical="center",
                    horizontal="center" if ci in (1, 4, 7) else "left")
                cell.font = font_ma if ci == 1 else font_data

        ws.freeze_panes = "A4"
        fname = f"DanhSach_NhanVien_{datetime.datetime.now().strftime('%d%m%Y_%H%M')}.xlsx"
        wb.save(fname)
        QMessageBox.information(self, "Excel", f"✅ Đã xuất: {fname}")

    def _xem_ho_so(self):
        if not self._check_sel("xem"): return
        row = self._get_row("nhan_vien", self._sel_id)
        if row:
            try:
                NhanVienHoSoDialog(self, row).exec()
            except Exception as e:
                print(f"❌ Lỗi mở hồ sơ: {e}")
                import traceback
                traceback.print_exc()
                QMessageBox.critical(self, "Lỗi", f"Không thể mở hồ sơ:\n{str(e)}")

    def _ai_cham_cong(self):
        if not self._check_sel("AI chấm công"): return
        row = self._get_row("nhan_vien", self._sel_id)
        if not row: return
        from PyQt6.QtWidgets import QInputDialog
        thang, ok = QInputDialog.getInt(self, "AI Chấm công",
                                        f"Sinh dữ liệu tháng mấy cho {row['ho_ten']}?",
                                        datetime.now().month, 1, 12)
        if not ok: return
        nam, ok2 = QInputDialog.getInt(self, "Năm", "Năm:",
                                       datetime.now().year, 2023, 2030)
        if not ok2: return
        AIChamCongWorker(self, row["id"], row["ho_ten"], thang, nam).run_and_show()

class NhanVienDialog(BaseDialog):
    def __init__(self,parent=None,data=None):
        super().__init__(parent,"Thêm/Sửa nhân viên",480); self.data=data
        self._add_title("🧑‍💼  THÔNG TIN NHÂN VIÊN"); form=self._add_form()
        self.f_ma=self._f_line("NV006"); self.f_ten=self._f_line("Họ tên đầy đủ")
        self.f_cv=self._f_combo(["Giám đốc","Quản lý BH","Nhân viên BH","Kỹ thuật viên","Kế toán"])
        self.f_sdt=self._f_line("0912 345 678"); self.f_email=self._f_line("nv@auto.vn")
        self.f_ngay=self._f_line("YYYY-MM-DD"); self.f_luong=self._f_spin(200e6,500000)
        self.f_tt=self._f_combo(["Đang làm","Thử việc","Nghỉ việc"])
        for l,w in [("Mã NV *",self.f_ma),("Họ tên *",self.f_ten),("Chức vụ",self.f_cv),
                    ("Số ĐT",self.f_sdt),("Email",self.f_email),("Ngày vào làm",self.f_ngay),
                    ("Lương",self.f_luong),("Trạng thái",self.f_tt)]: form.addRow(l,w)
        self._add_buttons("💾  Lưu NV")
        if data: self._fill(data)
    def _fill(self,d):
        self.f_ma.setText(d.get("ma_nv","")); self.f_ma.setReadOnly(True); self.f_ma.setStyleSheet("color:#6b7280;background:#13151c;")
        self.f_ten.setText(d.get("ho_ten",""))
        idx=self.f_cv.findText(d.get("chuc_vu","Nhân viên BH"))
        if idx>=0: self.f_cv.setCurrentIndex(idx)
        self.f_sdt.setText(d.get("so_dt","") or ""); self.f_email.setText(d.get("email","") or "")
        self.f_ngay.setText(d.get("ngay_vao","") or ""); self.f_luong.setValue(float(d.get("luong",0) or 0))
        idx2=self.f_tt.findText(d.get("trang_thai","Đang làm"))
        if idx2>=0: self.f_tt.setCurrentIndex(idx2)
    def _save(self):
        ma=self.f_ma.text().strip(); ten=self.f_ten.text().strip()
        if not all([ma,ten]): QMessageBox.warning(self,"","Điền đủ Mã NV và Họ tên!"); return
        conn=get_conn()
        try:
            vals=(ten,self.f_cv.currentText(),self.f_sdt.text(),self.f_email.text(),self.f_ngay.text(),self.f_luong.value(),self.f_tt.currentText())
            if self.data: conn.execute("UPDATE nhan_vien SET ho_ten=?,chuc_vu=?,so_dt=?,email=?,ngay_vao=?,luong=?,trang_thai=? WHERE id=?",vals+(self.data["id"],))
            else: conn.execute("INSERT INTO nhan_vien(ma_nv,ho_ten,chuc_vu,so_dt,email,ngay_vao,luong,trang_thai) VALUES(?,?,?,?,?,?,?,?)",(ma,)+vals)
            conn.commit(); self.accept()
        except Exception as e: QMessageBox.critical(self,"Lỗi",str(e))
        finally: conn.close()

# ════════════════════════════════════════════════════════════════
# AI CHẤM CÔNG + HỒ SƠ NHÂN VIÊN
# ════════════════════════════════════════════════════════════════
class AIChamCongWorker:
    def __init__(self, parent, nv_id, ho_ten, thang, nam):
        self.parent = parent; self.nv_id = nv_id
        self.ho_ten = ho_ten; self.thang = thang; self.nam = nam

    def run_and_show(self):
        from PyQt6.QtWidgets import QProgressDialog
        from PyQt6.QtCore import Qt
        prog = QProgressDialog(f"🤖 AI đang sinh dữ liệu cho {self.ho_ten}...",
                               None, 0, 0, self.parent)
        prog.setWindowTitle("AI Chấm công")
        prog.setWindowModality(Qt.WindowModality.WindowModal)
        prog.show()
        try:
            days_in_month = calendar.monthrange(self.nam, self.thang)[1]
            conn = get_conn()
            conn.execute("DELETE FROM cham_cong WHERE nv_id=? AND strftime('%Y-%m',ngay)=?",
                (self.nv_id, f"{self.nam}-{self.thang:02d}"))
            records = []
            for d in range(1, days_in_month+1):
                day = date(self.nam, self.thang, d)
                if day.weekday() >= 5 or day > date.today(): continue
                d_str = day.strftime("%Y-%m-%d")
                rand = random.random()
                if rand < 0.02:
                    records.append((self.nv_id,d_str,None,None,"Vắng mặt",200000,0,"AI: Vắng không lý do"))
                elif rand < 0.05:
                    gv = f"{random.randint(8,9):02d}:{random.randint(20,59):02d}"
                    gr = f"{random.randint(12,14):02d}:{random.randint(0,59):02d}"
                    records.append((self.nv_id,d_str,gv,gr,"Nửa buổi",100000,0,"AI: Làm nửa buổi"))
                elif rand < 0.10:
                    gv = f"08:{random.randint(16,45):02d}"
                    gr = f"{random.randint(17,18):02d}:{random.randint(30,59):02d}"
                    records.append((self.nv_id,d_str,gv,gr,"Đi muộn",50000,0,f"AI: Đến muộn {gv}"))
                elif rand < 0.13:
                    gv = f"07:{random.randint(45,59):02d}"
                    gr = f"16:{random.randint(30,59):02d}"
                    records.append((self.nv_id,d_str,gv,gr,"Về sớm",50000,0,f"AI: Về sớm {gr}"))
                else:
                    gv = f"07:{random.randint(45,59):02d}"
                    gr = f"{random.randint(17,18):02d}:{random.randint(30,59):02d}"
                    thuong = 100000 if gr >= "18:00" else 0
                    gc = "AI: Làm thêm giờ" if thuong > 0 else "AI: Đúng giờ"
                    records.append((self.nv_id,d_str,gv,gr,"Đúng giờ",0,thuong,gc))
            conn.executemany("""INSERT INTO cham_cong
                (nv_id,ngay,gio_vao,gio_ra,trang_thai,phat,thuong,ghi_chu)
                VALUES(?,?,?,?,?,?,?,?)""", records)
            conn.commit(); conn.close(); prog.close()
            dung_gio = sum(1 for r in records if r[4]=="Đúng giờ")
            di_muon  = sum(1 for r in records if r[4]=="Đi muộn")
            vang     = sum(1 for r in records if r[4]=="Vắng mặt")
            QMessageBox.information(self.parent, "✅ AI Chấm công xong!",
                f"🤖 Đã sinh {len(records)} ngày công cho {self.ho_ten}\n"
                f"Tháng {self.thang}/{self.nam}\n\n"
                f"✅ Đúng giờ: {dung_gio} ngày\n"
                f"⏰ Đi muộn: {di_muon} ngày\n"
                f"❌ Vắng mặt: {vang} ngày")
        except Exception as e:
            prog.close(); QMessageBox.critical(self.parent,"Lỗi AI",str(e))

class NhanVienHoSoDialog(QDialog):
    def __init__(self, parent=None, data=None):
        super().__init__(parent)
        self.data = data or {}; self.nv_id = data.get("id") if data else None
        self.setWindowTitle(f"Hồ sơ — {data.get('ho_ten','')}")
        screen = QApplication.primaryScreen().availableGeometry()
        self.setMinimumSize(min(820, int(screen.width()*0.85)), min(620, int(screen.height()*0.82)))
        self.setStyleSheet("""
            QDialog{background:#f0f4f8;color:#00274c;}
            QLabel{color:#1e40af;font-size:13px;font-weight:600;background:transparent;letter-spacing:0.3px;}
            QLineEdit,QTextEdit,QDoubleSpinBox,QComboBox{
                background:#ffffff;color:#00274c;border:0.5px solid #dbeafe;
                border-radius:8px;padding:8px 12px;font-size:13px;}
            QLineEdit:focus,QTextEdit:focus,QDoubleSpinBox:focus,QComboBox:focus{border-color:#2563eb;}
            QComboBox QAbstractItemView{background:#ffffff;color:#00274c;
                border:0.5px solid #dbeafe;selection-background-color:#eff6ff;}
            QPushButton{background:#ffffff;color:#00274c;border:0.5px solid #dbeafe;
                border-radius:8px;padding:8px 18px;font-size:13px;}
            QPushButton:hover{background:#eff6ff;color:#1e40af;}
            QPushButton#btn_save{background:#2563eb;color:white;border:none;font-weight:700;min-width:100px;}
            QPushButton#btn_save:hover{background:#1d4ed8;}
            QPushButton#btn_pdf{background:#eff6ff;color:#1e40af;border:0.5px solid #bfdbfe;font-weight:600;}
            QPushButton#btn_pdf:hover{background:#2563eb;color:white;}
            QTabWidget::pane{border:0.5px solid #dbeafe;border-radius:8px;}
            QTabBar::tab{padding:6px 16px;font-size:12px;font-weight:600;
                color:#64748b;background:#f8fafc;border:none;
                border-bottom:2px solid transparent;}
            QTabBar::tab:selected{color:#2563eb;border-bottom:2px solid #2563eb;}
        """)
        self._build(); self._load()

    def _build(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        # ── HEADER ───────────────────────────────────────────────────────
        hdr = QWidget()
        hdr.setStyleSheet("background:#ffffff;border-bottom:3px solid #2563eb;padding:20px;")
        hdr_l = QVBoxLayout(hdr)
        hdr_l.setContentsMargins(0, 0, 0, 0)
        hdr_l.setSpacing(16)

        # Row 1: Avatar + Tên + Trạng thái
        row1 = QHBoxLayout()
        av = QLabel("🧑‍💼")
        av.setStyleSheet("font-size:64px;")
        av.setFixedSize(100, 100)
        av.setAlignment(Qt.AlignmentFlag.AlignCenter)
        row1.addWidget(av)

        col_ten = QVBoxLayout()
        col_ten.setSpacing(4)
        ten = QLabel(self.data.get("ho_ten", "").upper())
        ten.setStyleSheet("font-size:26px;font-weight:900;color:#111827;background:transparent;")
        ma_lbl = QLabel(f"🪪 {self.data.get('ma_nv', '')}")
        ma_lbl.setStyleSheet("font-size:14px;color:#0284c7;font-weight:700;background:transparent;")
        cv_lbl = QLabel(f"💼 {self.data.get('chuc_vu', '')}")
        cv_lbl.setStyleSheet("font-size:14px;color:#0284c7;font-weight:700;background:transparent;")
        col_ten.addWidget(ten)
        col_ten.addWidget(ma_lbl)
        col_ten.addWidget(cv_lbl)
        row1.addLayout(col_ten, 1)

        # Bên phải: Trạng thái + Lương
        col_right = QVBoxLayout()
        col_right.setSpacing(8)
        tt = self.data.get("trang_thai", "")
        tt_color = {"Đang làm": "#10b981", "Thử việc": "#f59e0b", "Nghỉ việc": "#ef4444"}.get(tt, "#6b7280")
        tt_lbl = QLabel(f"● {tt}")
        tt_lbl.setStyleSheet(f"font-size:16px;font-weight:900;color:{tt_color};background:transparent;text-align:right;")
        luong = int(self.data.get('luong', 0) or 0)
        luong_lbl = QLabel(f"💰 {luong:,}₫/tháng")
        luong_lbl.setStyleSheet("font-size:14px;font-weight:800;color:#059669;background:transparent;text-align:right;")
        col_right.addWidget(tt_lbl)
        col_right.addWidget(luong_lbl)
        row1.addLayout(col_right)
        hdr_l.addLayout(row1)

        # Row 2: Contact info
        row2 = QHBoxLayout()
        row2.setContentsMargins(100, 0, 0, 0)
        sdt = QLabel(f"📱 {self.data.get('so_dt', '') or '—'}")
        sdt.setStyleSheet("font-size:13px;color:#d97706;font-weight:700;background:transparent;")
        email = QLabel(f"📧 {self.data.get('email', '') or '—'}")
        email.setStyleSheet("font-size:13px;color:#d97706;font-weight:700;background:transparent;")
        ngay = QLabel(f"📅 Vào: {self.data.get('ngay_vao', '') or '—'}")
        ngay.setStyleSheet("font-size:13px;color:#059669;font-weight:700;background:transparent;")
        row2.addWidget(sdt)
        row2.addWidget(email)
        row2.addWidget(ngay)
        row2.addStretch()
        hdr_l.addLayout(row2)

        root.addWidget(hdr)

        # ── STAT CARDS ───────────────────────────────────────────────────
        stat_w = QWidget()
        stat_w.setStyleSheet("background:#f9fafb;border-bottom:1px solid #e5e7eb;")
        stat_l = QHBoxLayout(stat_w)
        stat_l.setContentsMargins(16, 14, 16, 14)
        stat_l.setSpacing(12)

        self._stat_lbls = {}
        for key, icon, label, color in [
            ("di_lam", "✅", "Ngày công", "#0F1F35"),
            ("di_muon", "⏰", "Đi muộn", "#059669"),
            ("vang", "❌", "Vắng", "#7c3aed"),
            ("don_hang", "📋", "Đơn hàng", "#d97706"),
        ]:
            c = QWidget()
            c.setStyleSheet(
                f"background:#ffffff;border:1px solid #e5e7eb;"
                f"border-radius:12px;border-top:3px solid {color};")
            cl = QVBoxLayout(c)
            cl.setContentsMargins(14, 12, 14, 12)
            cl.setSpacing(2)
            val_lbl = QLabel("—")
            val_lbl.setStyleSheet(
                f"font-size:18px;font-weight:900;color:{color};background:transparent;")
            lbl_lbl = QLabel(f"{icon} {label}")
            lbl_lbl.setStyleSheet("font-size:12px;color:#6b7280;font-weight:700;background:transparent;")
            cl.addWidget(val_lbl)
            cl.addWidget(lbl_lbl)
            stat_l.addWidget(c, 1)
            self._stat_lbls[key] = val_lbl

        root.addWidget(stat_w)

        # ── NỘI DUNG CHÍNH ───────────────────────────────────────────────
        content = QWidget()
        content.setStyleSheet("background:#ffffff;")
        content_l = QVBoxLayout(content)
        content_l.setContentsMargins(24, 20, 24, 20)
        content_l.setSpacing(20)

        tabs = QTabWidget()

        # Tab chấm công
        t1 = QWidget()
        t1.setStyleSheet("background:#ffffff;")
        t1l = QVBoxLayout(t1)
        t1l.setContentsMargins(0, 0, 0, 0)

        cols_cc = ["NGÀY", "THỨ", "GIỜ VÀO", "GIỜ RA", "TRẠNG THÁI", "PHẠT", "THƯỞNG", "GHI CHÚ"]
        self.tbl_cc = QTableWidget(0, len(cols_cc))
        self.tbl_cc.setHorizontalHeaderLabels(cols_cc)
        self.tbl_cc.setAlternatingRowColors(True)
        self.tbl_cc.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.tbl_cc.setShowGrid(False)
        self.tbl_cc.verticalHeader().setVisible(False)
        self.tbl_cc.horizontalHeader().setSectionResizeMode(7, QHeaderView.ResizeMode.Stretch)
        t1l.addWidget(self.tbl_cc)
        tabs.addTab(t1, "⏰  Chấm công tháng này")

        # Tab đơn hàng
        t2 = QWidget()
        t2.setStyleSheet("background:#ffffff;")
        t2l = QVBoxLayout(t2)
        t2l.setContentsMargins(0, 0, 0, 0)

        cols_dh = ["MÃ ĐƠN", "XE", "KHÁCH HÀNG", "GIÁ BÁN", "TRẠNG THÁI", "NGÀY ĐẶT"]
        self.tbl_dh = QTableWidget(0, len(cols_dh))
        self.tbl_dh.setHorizontalHeaderLabels(cols_dh)
        self.tbl_dh.setAlternatingRowColors(True)
        self.tbl_dh.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.tbl_dh.setShowGrid(False)
        self.tbl_dh.verticalHeader().setVisible(False)
        self.tbl_dh.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        t2l.addWidget(self.tbl_dh)
        tabs.addTab(t2, "📋  Đơn hàng")

        content_l.addWidget(tabs, 1)
        root.addWidget(content, 1)

        # ── FOOTER ───────────────────────────────────────────────────────
        ftr = QWidget()
        ftr.setFixedHeight(60)
        ftr.setStyleSheet("background:#ffffff;border-top:1px solid #e5e7eb;")
        fl = QHBoxLayout(ftr)
        fl.setContentsMargins(24, 14, 24, 14)
        fl.setSpacing(12)
        fl.addStretch()

        btn_close = QPushButton("✖  Đóng")
        btn_close.setObjectName("btn_close")
        btn_close.setMinimumWidth(120)
        btn_close.setMinimumHeight(40)
        btn_close.setCursor(Qt.CursorShape.PointingHandCursor)
        btn_close.clicked.connect(self.reject)

        fl.addWidget(btn_close)
        root.addWidget(ftr)

    def _load(self):
        if not self.nv_id: return
        now = datetime.now(); thang=now.month; nam=now.year
        conn = get_conn()
        cc_rows = conn.execute("""SELECT ngay,gio_vao,gio_ra,trang_thai,phat,thuong,ghi_chu
            FROM cham_cong WHERE nv_id=? AND strftime('%Y-%m',ngay)=?
            ORDER BY ngay DESC""", (self.nv_id,f"{nam}-{thang:02d}")).fetchall()
        dh_rows = conn.execute("""SELECT dh.ma_don,x.hang_xe||' '||x.dong_xe,kh.ho_ten,
            dh.gia_ban_thuc,dh.trang_thai,dh.ngay_dat
            FROM don_hang dh JOIN xe x ON dh.xe_id=x.id
            JOIN khach_hang kh ON dh.kh_id=kh.id
            WHERE dh.nv_id=? ORDER BY dh.id DESC""", (self.nv_id,)).fetchall()
        nv = conn.execute("SELECT luong FROM nhan_vien WHERE id=?",(self.nv_id,)).fetchone()
        conn.close()
        luong_cb = float(nv[0] or 0) if nv else 0
        days_in = calendar.monthrange(nam,thang)[1]
        work_days = sum(1 for d in range(1,days_in+1) if date(nam,thang,d).weekday()<5)
        di_lam  = sum(1 for r in cc_rows if r[3]!="Vắng mặt")
        di_muon = sum(1 for r in cc_rows if r[3]=="Đi muộn")
        vang    = sum(1 for r in cc_rows if r[3]=="Vắng mặt")
        tong_phat   = sum(float(r[4] or 0) for r in cc_rows)
        tong_thuong = sum(float(r[5] or 0) for r in cc_rows)
        luong_tt = (luong_cb/work_days*di_lam if work_days>0 else 0)-tong_phat+tong_thuong
        doanh_so = sum(r[3] for r in dh_rows)
        self._stat_lbls["di_lam"].setText(f"{di_lam}/{work_days}")
        self._stat_lbls["di_muon"].setText(str(di_muon))
        self._stat_lbls["vang"].setText(str(vang))
        self._stat_lbls["don_hang"].setText(str(len(dh_rows)))
        # Bảng chấm công
        thu_map={0:"Thứ 2",1:"Thứ 3",2:"Thứ 4",3:"Thứ 5",4:"Thứ 6",5:"Thứ 7",6:"CN"}
        TT_COLOR={"Đúng giờ":"#4ade80","Đi muộn":"#fbbf24","Về sớm":"#f97316",
                  "Vắng mặt":"#f87171","Nửa buổi":"#a78bfa","Nghỉ phép":"#60a5fa"}
        self.tbl_cc.setRowCount(0)
        for row in cc_rows:
            r=self.tbl_cc.rowCount(); self.tbl_cc.insertRow(r); self.tbl_cc.setRowHeight(r,40)
            try: d=datetime.strptime(row[0],"%Y-%m-%d"); ngay_fmt=d.strftime("%d/%m/%Y"); thu=thu_map[d.weekday()]
            except Exception: ngay_fmt=row[0]; thu=""
            tt=row[3] or "—"; color=TT_COLOR.get(tt,"#94a3b8")
            phat=float(row[4] or 0); thuong=float(row[5] or 0)
            for c,val in enumerate([ngay_fmt,thu,row[1] or "—",row[2] or "—",tt,
                f"-{phat/1000:.0f}k" if phat>0 else "—",
                f"+{thuong/1000:.0f}k" if thuong>0 else "—",row[6] or ""]):
                item=QTableWidgetItem(val)
                if c==4: item.setForeground(QColor(color)); item.setFont(QFont("Segoe UI",11,QFont.Weight.Bold))
                elif c==5 and phat>0: item.setForeground(QColor("#f87171"))
                elif c==6 and thuong>0: item.setForeground(QColor("#4ade80"))
                self.tbl_cc.setItem(r,c,item)
        # Bảng đơn hàng
        STATUS_COL={"Đã giao xe":"#4ade80","Đã thanh toán":"#60a5fa","Chờ xử lý":"#fbbf24"}
        self.tbl_dh.setRowCount(0)
        for row in dh_rows:
            r=self.tbl_dh.rowCount(); self.tbl_dh.insertRow(r); self.tbl_dh.setRowHeight(r,40)
            for c,val in enumerate([row[0],row[1],row[2],f"{row[3]/1e9:.2f} tỷ",row[4],row[5] or ""]):
                item=QTableWidgetItem(val)
                if c==0: item.setForeground(QColor("#a78bfa")); item.setFont(QFont("Segoe UI",11,QFont.Weight.Bold))
                elif c==3: item.setForeground(QColor("#4ade80")); item.setFont(QFont("Segoe UI",12,QFont.Weight.Bold))
                elif c==4: item.setForeground(QColor(STATUS_COL.get(val,"#94a3b8")))
                self.tbl_dh.setItem(r,c,item)
